import json, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
HERE = os.path.dirname(os.path.abspath(__file__)); SEED = os.path.join(HERE, 'grc', 'seed_data')
D = json.load(open(os.path.join(SEED, 'stage4_all_domains.json'), encoding='utf-8'))['domains']
s3 = json.load(open(os.path.join(SEED, 'stage3_complete.json'), encoding='utf-8'))
def fwshort(c): return c['framework'].split('(')[0].strip()
def sfn(f):
    for k, v in [('22301', 'ISO 22301'), ('42001', 'ISO 42001'), ('27001', 'ISO 27001'), ('Artificial', 'NIST AI RMF'), ('800-53', 'NIST 800-53'), ('Cybersecurity Framework', 'NIST CSF'), ('Health Information', 'DOH ADHIE'), ('ETGRMF', 'SBP ETGRMF'), ('Internet Banking', 'SBP IB'), ('SBP Cloud', 'SBP Cloud'), ('Personal Data Transfer', 'KSA Transfer'), ('Abu Dhabi', 'ADHICS'), ('General Data', 'GDPR'), ('National Data', 'KSA NDMO'), ('Digital Operational', 'DORA')]:
        if k in f: return v
    return f.split('(')[0].strip()
raw = {}
for dm in s3['domains']:
    for c in dm['controls']:
        raw[(fwshort(c), c['control_id'], c['title'])] = {'ev': [e['name'] for e in (c.get('evidence') or [])], 'art': [a.get('name') for a in (c.get('artifacts') or [])]}
hd = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='17B898'); thin = Side(style='thin', color='DDDDDD'); band = PatternFill('solid', fgColor='E8FCF8')
def hdr(ws, heads):
    for j, h in enumerate(heads, 1):
        c = ws.cell(1, j, h); c.font = hd; c.fill = fill; c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.freeze_panes = 'A2'
def save(wb, name):
    xp = os.path.expanduser('~/OneDrive/Desktop/' + name + '.xlsx')
    try: wb.save(xp)
    except PermissionError: xp = os.path.expanduser('~/OneDrive/Desktop/' + name + '_v2.xlsx'); wb.save(xp)
    return xp

# ===== MAIN EXCEL =====
wb = openpyxl.Workbook()
sm = wb.active; sm.title = 'Summary'
hdr(sm, ['Domain', 'Controls', 'Normalized sets', 'Standalone', '# Frameworks', '# Absent frameworks'])
r = 2
for dm in D:
    multi = [s for s in dm['sets'] if s['member_count'] > 1]; single = [s for s in dm['sets'] if s['member_count'] == 1]
    for col, val in zip(range(1, 7), [dm['domain'], dm['controls_in'], len(multi), len(single), len(dm['frameworks']), len(dm.get('absent_frameworks', []))]):
        sm.cell(r, col, val).border = Border(bottom=thin)
    r += 1
for col, w in zip('ABCDEF', [40, 10, 15, 11, 13, 18]): sm.column_dimensions[col].width = w

ns = wb.create_sheet('Normalized sets (full)')
hdr(ns, ['Domain', 'Set', 'Heading', 'Framework', 'Control ID', 'Original Title', 'Normalized evidence (set)', 'Excluded evidence + reason', 'Requirement-specific artifacts'])
r = 2; flip = False
for dm in D:
    for s in [x for x in dm['sets'] if x['member_count'] > 1]:
        start = r; flip = not flip
        nev = '\n'.join('* ' + e['name'] for e in s['normalized_evidence']) or '(none)'
        exc = '\n'.join('- ' + x['name'] + ' :: ' + x.get('reason', '') for x in s.get('excluded_evidence', [])) or '(none)'
        nart = '\n'.join('* ' + a['name'] for a in s['normalized_artifacts']) or '(none - see framework-catalog tab)'
        for m in s['members']:
            ns.cell(r, 4, sfn(m['framework'])); ns.cell(r, 5, m['control_id']); ns.cell(r, 6, m['original_title']); r += 1
        end = r - 1
        for col, val in ((1, dm['domain']), (2, s['set_id']), (3, s['normalized_title']), (7, nev), (8, exc), (9, nart)):
            if end > start: ns.merge_cells(start_row=start, start_column=col, end_row=end, end_column=col)
            ns.cell(start, col, val).alignment = Alignment(vertical='top', wrap_text=True)
        for rr in range(start, end + 1):
            for col in range(1, 10):
                ns.cell(rr, col).alignment = Alignment(vertical='top', wrap_text=True); ns.cell(rr, col).border = Border(bottom=thin)
                if flip: ns.cell(rr, col).fill = band
for col, w in zip('ABCDEFGHI', [22, 6, 32, 14, 11, 34, 46, 40, 38]): ns.column_dimensions[col].width = w

af = wb.create_sheet('Absent frameworks (why)')
hdr(af, ['Domain', 'Absent framework', 'Reason'])
r = 2
for dm in D:
    for a in dm.get('absent_frameworks', []):
        af.cell(r, 1, dm['domain']); af.cell(r, 2, sfn(a.get('name', ''))); af.cell(r, 3, a.get('reason', ''))
        for col in range(1, 4):
            af.cell(r, col).alignment = Alignment(vertical='top', wrap_text=True); af.cell(r, col).border = Border(bottom=thin)
        r += 1
for col, w in zip('ABC', [36, 22, 90]): af.column_dimensions[col].width = w

fc = wb.create_sheet('Framework-catalog artifacts')
hdr(fc, ['Domain', 'Framework', 'Note', 'Catalog artifacts (not requirement-specific)'])
r = 2
for dm in D:
    for cat in dm.get('framework_catalog_artifacts', []):
        fc.cell(r, 1, dm['domain']); fc.cell(r, 2, sfn(cat['framework'])); fc.cell(r, 3, cat['note'])
        fc.cell(r, 4, '\n'.join('* ' + a['name'] for a in cat['artifacts']))
        for col in range(1, 5):
            fc.cell(r, col).alignment = Alignment(vertical='top', wrap_text=True); fc.cell(r, col).border = Border(bottom=thin)
        r += 1
for col, w in zip('ABCD', [30, 16, 55, 50]): fc.column_dimensions[col].width = w

# Standalone controls (full) — every framework-unique control, individually listed
st = wb.create_sheet('Standalone controls (full)')
hdr(st, ['Domain', 'Framework', 'Control ID', 'Original Title'])
r = 2; nstand = 0
for dm in D:
    for s in [x for x in dm['sets'] if x['member_count'] == 1]:
        m = s['members'][0]; nstand += 1
        st.cell(r, 1, dm['domain']); st.cell(r, 2, sfn(m['framework'])); st.cell(r, 3, m['control_id']); st.cell(r, 4, m['original_title'])
        for col in range(1, 5):
            st.cell(r, col).alignment = Alignment(vertical='top', wrap_text=True); st.cell(r, col).border = Border(bottom=thin)
        r += 1
for col, w in zip('ABCD', [36, 18, 14, 70]): st.column_dimensions[col].width = w
print('standalone controls listed:', nstand)
print('MAIN Excel:', save(wb, 'UNIFIED_LIBRARY_all_domains'))

# ===== RAW vs NORMALIZED =====
wb2 = openpyxl.Workbook(); ws = wb2.active; ws.title = 'Raw vs Normalized'
hdr(ws, ['Domain', 'Set', 'Normalized Heading', 'Framework', 'Control ID', 'Original Title', 'RAW evidence (this control)', 'RAW artifacts', 'NORMALIZED evidence (set) -- name <= raw absorbed', 'Excluded (reason)', 'Requirement-specific artifacts'])
r = 2; flip = False; nsets = 0
for dm in D:
    for s in [x for x in dm['sets'] if x['member_count'] > 1]:
        nsets += 1; start = r; flip = not flip
        nev = '\n'.join('* ' + e['name'] + '  <= [' + '; '.join(e.get('absorbs', [])) + ']' for e in s['normalized_evidence']) or '(none)'
        exc = '\n'.join('- ' + x['name'] + ' :: ' + x.get('reason', '') for x in s.get('excluded_evidence', [])) or '(none)'
        nart = '\n'.join('* ' + a['name'] + ' [' + str(a.get('type')) + ']' for a in s['normalized_artifacts']) or '(none - framework-catalog only)'
        for m in s['members']:
            rr = raw.get((m['framework'], m['control_id'], m['original_title']), {'ev': [], 'art': []})
            ws.cell(r, 4, sfn(m['framework'])); ws.cell(r, 5, m['control_id']); ws.cell(r, 6, m['original_title'])
            ws.cell(r, 7, '\n'.join('- ' + x for x in rr['ev']) or '(none)'); ws.cell(r, 8, '\n'.join('- ' + str(x) for x in rr['art']) or '(none)')
            r += 1
        end = r - 1
        for col, val in ((1, dm['domain']), (2, s['set_id']), (3, s['normalized_title']), (9, nev), (10, exc), (11, nart)):
            if end > start: ws.merge_cells(start_row=start, start_column=col, end_row=end, end_column=col)
            ws.cell(start, col, val).alignment = Alignment(vertical='top', wrap_text=True)
        for rr2 in range(start, end + 1):
            for col in range(1, 12):
                ws.cell(rr2, col).alignment = Alignment(vertical='top', wrap_text=True); ws.cell(rr2, col).border = Border(bottom=thin)
                if flip: ws.cell(rr2, col).fill = band
for col, w in zip('ABCDEFGHIJK', [20, 6, 28, 13, 10, 30, 34, 24, 48, 30, 30]): ws.column_dimensions[col].width = w
print('RAW-vs-NORMALIZED Excel (%d sets):' % nsets, save(wb2, 'UNIFIED_LIBRARY_raw_vs_normalized'))
