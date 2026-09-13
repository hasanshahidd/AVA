import json, os, re
from collections import Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

HERE = os.path.dirname(__file__); SEED = os.path.join(HERE, 'grc', 'seed_data')
d = json.load(open(os.path.join(SEED, 'stage4_sets_governance.json'), encoding='utf-8'))
s3 = json.load(open(os.path.join(SEED, 'stage3_complete.json'), encoding='utf-8'))
gov = [x for x in s3['domains'] if x['domain'].startswith('Governance')][0]['controls']
def fwshort(c): return c['framework'].split('(')[0].strip()
def nk(s): return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9 ]', ' ', (s or '').lower())).strip()
rawev = {}
for c in gov: rawev[(fwshort(c), c['control_id'], c['title'])] = [{'name': e['name']} for e in (c.get('evidence') or [])]

audit = {}
for fn in ['_gov_audit_00_13.json', '_gov_audit_14_27.json', '_gov_audit_28_40.json']:
    for r in json.load(open(os.path.join(HERE, fn), encoding='utf-8'))['results']:
        audit[r['set_id']] = r

S = {s['set_id']: s for s in d['sets']}
new_sets = []; split_count = 0; ev_dropped = 0
for s in d['sets']:
    a = audit.get(s['set_id'])
    if not a or a.get('verdict') == 'ACCURATE':
        new_sets.append(s); continue
    loose = {(x['fw'], x['cid']) for x in (a.get('loose_members') or [])}
    offtopic = {nk(x) for x in (a.get('offtopic_evidence') or [])}
    kept = [m for m in s['members'] if (fwshort_m := m['framework'].split('(')[0].strip(), m['control_id']) and (m['framework'].split('(')[0].strip(), m['control_id']) not in loose]
    out = [m for m in s['members'] if (m['framework'].split('(')[0].strip(), m['control_id']) in loose]
    # drop off-topic evidence from the set
    before = len(s['normalized_evidence'])
    ne = [e for e in s['normalized_evidence'] if nk(e['name']) not in offtopic]
    ev_dropped += before - len(ne)
    if len(kept) >= 2:
        s2 = dict(s); s2['members'] = kept; s2['member_count'] = len(kept)
        s2['frameworks'] = sorted(set(m['framework'].split('(')[0].strip() for m in kept)); s2['normalized_evidence'] = ne
        new_sets.append(s2)
    else:
        # set collapsed to <2 -> every member becomes standalone
        out = s['members']
    for m in out:
        split_count += 1
        own = rawev.get((m['framework'].split('(')[0].strip(), m['control_id'], m['original_title']), [])
        seen = {}; oe = []
        for e in own:
            k = nk(e['name'])
            if k and k not in seen: seen[k] = 1; oe.append({'name': e['name'], 'absorbs': [e['name']], 'sources': [m['framework'].split('(')[0].strip() + ' ' + m['control_id']]})
        new_sets.append({'set_id': s['set_id'] + '-x', 'normalized_title': m['original_title'], 'member_count': 1,
                         'frameworks': [m['framework'].split('(')[0].strip()], 'members': [m],
                         'normalized_evidence': oe, 'normalized_artifacts': []})

# verify
allfw_ok = all(len(set(m['framework'].split('(')[0].strip() for m in s['members'])) == s['member_count'] for s in new_sets)
totctrl = sum(s['member_count'] for s in new_sets)
multi = [s for s in new_sets if s['member_count'] > 1]; single = [s for s in new_sets if s['member_count'] == 1]
import statistics
evc = [len(s['normalized_evidence']) for s in multi]
print('controls total:', totctrl, '(expect 420) | no-same-framework:', 'OK' if allfw_ok else 'FAIL')
print('loose members split out:', split_count, '| off-topic evidence items dropped:', ev_dropped)
print('sets: %d normalized + %d standalone' % (len(multi), len(single)))
print('normalized-set evidence: median', int(statistics.median(evc)), 'max', max(evc), '| >8:', sum(1 for x in evc if x > 8))

# renumber
new_sets.sort(key=lambda s: (-s['member_count']))
for j, s in enumerate(new_sets, 1): s['set_id'] = 'GOV-%03d' % j
d['sets'] = new_sets; d['normalized_sets'] = len(multi); d['standalone'] = len(single); d['sets_out'] = len(new_sets)
d['evidence_normalized'] = sum(len(s['normalized_evidence']) for s in new_sets)
json.dump(d, open(os.path.join(SEED, 'stage4_sets_governance.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

# regenerate combined public data
allj = json.load(open(os.path.join(SEED, 'stage3_complete.json'), encoding='utf-8'))
doms = []
for dd in allj['domains']:
    rec = {'name': dd['domain'], 'controls': len(dd['controls']), 'frameworks': len(set(fwshort(c) for c in dd['controls'])), 'status': 'pending', 'normalized_sets': 0, 'standalone': 0}
    if dd['domain'].startswith('Governance'): rec['status'] = 'normalized'; rec['normalized_sets'] = d['normalized_sets']; rec['standalone'] = d['standalone']
    doms.append(rec)
open(os.path.join(HERE, '..', 'grc-frontend', 'public', 'library_data.json'), 'w', encoding='utf-8').write(json.dumps({'domains': doms, 'governance': d}, ensure_ascii=False))

wb = openpyxl.Workbook(); hd = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='2E7D6B'); thin = Side(style='thin', color='DDDDDD')
sm = wb.active; sm.title = 'Sets summary'
for j, h in enumerate(['Set ID', 'Normalized Heading', '# Frameworks', '# Evidence', '# Artifacts', 'Type'], 1):
    cc = sm.cell(1, j, h); cc.font = hd; cc.fill = fill
r = 2
for s in new_sets:
    sm.cell(r, 1, s['set_id']); sm.cell(r, 2, s['normalized_title']); sm.cell(r, 3, s['member_count']); sm.cell(r, 4, len(s['normalized_evidence'])); sm.cell(r, 5, len(s['normalized_artifacts'])); sm.cell(r, 6, 'standalone' if s['member_count'] == 1 else 'set')
    for col in range(1, 7): sm.cell(r, col).border = Border(bottom=thin)
    r += 1
sm.freeze_panes = 'A2'
for col, w in zip('ABCDEF', [12, 56, 12, 10, 10, 12]): sm.column_dimensions[col].width = w
xp = os.path.expanduser('~/OneDrive/Desktop/GOVERNANCE_normalized.xlsx')
try: wb.save(xp)
except PermissionError: xp = os.path.expanduser('~/OneDrive/Desktop/GOVERNANCE_normalized_v2.xlsx'); wb.save(xp)
print('updated JSON + preview data + EXCEL:', xp)
