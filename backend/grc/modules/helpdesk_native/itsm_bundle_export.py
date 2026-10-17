"""ITSM workbook -> versioned JSON bundle exporter.

The owner's ITIL-4 ITSM workbook (AVA_ServiceDesk_ITSM_enriched_v*.xlsx) is the
AUTHORING tool; this turns it into ONE deterministic JSON bundle that
`itsm_import.import_bundle` ingests per tenant.

    python -m grc.modules.helpdesk_native.itsm_bundle_export <xlsx> <out.json>
    python -m grc.modules.helpdesk_native.itsm_bundle_export <out.json>   # auto-pick newest workbook

Bundle shape:
    {"bundle_version": "...", "generated_at": iso, "source_filename": "...",
     "checksum": sha256(entities), "entities": {<entity_key>: [ {col: val}, ... ]}}

* entity_key per sheet is read from the workbook's "Index" sheet (column
  "Entity key (JSON)", minus ".json"), e.g. business_services, sample_cis.
* Sheets are matched to Index rows by their numeric prefix ("02", "13", ...) so
  the dash character in the sheet titles (en/em dash, mojibake-prone) never has
  to be reproduced.
* Column headers are normalised to snake_case ("Supports Business Service Id"
  -> supports_business_service_id). Values are stripped; blank -> None.
* Row order is the sheet's own order => deterministic; the checksum is over the
  canonical JSON of `entities` only (so generated_at never perturbs it).

openpyxl is imported lazily so the importer (which only needs
`compute_checksum`) does not depend on it.
"""
from __future__ import annotations

import argparse
import datetime as _dt
import hashlib
import json
import os
import re
import sys
from typing import Any

DOWNLOADS = os.path.join(os.path.expanduser("~"), "Downloads")
# Preference order when no workbook is named on the command line.
_CANDIDATES = (
    "AVA_ServiceDesk_ITSM_enriched_v3.xlsx",
    "AVA_ServiceDesk_ITSM_enriched_v2.xlsx",
    "AVA_ServiceDesk_ITSM_enriched.xlsx",
    "AVA_ServiceDesk_ITSM.xlsx",
)
_INDEX_SHEET = "Index"


# --- helpers ------------------------------------------------------------------
def snake(header: Any) -> str:
    """'Supports Business Service Id' -> 'supports_business_service_id'."""
    s = re.sub(r"[^0-9A-Za-z]+", "_", str(header).strip()).strip("_").lower()
    return s


def _clean(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    if isinstance(v, bool):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.isoformat()
    return v


def compute_checksum(entities: dict) -> str:
    """sha256 of the canonical JSON of the entities map (shared with the importer
    so it can detect a bundle edited after export)."""
    blob = json.dumps(entities, sort_keys=True, ensure_ascii=False,
                      separators=(",", ":"), default=str)
    return hashlib.sha256(blob.encode("utf-8")).hexdigest()


def pick_workbook(explicit: str | None = None, search_dir: str = DOWNLOADS) -> str:
    """Explicit path wins; else the newest of v3 > v2 > enriched > base."""
    if explicit:
        if not os.path.isfile(explicit):
            raise FileNotFoundError(explicit)
        return explicit
    for name in _CANDIDATES:
        p = os.path.join(search_dir, name)
        if os.path.isfile(p):
            return p
    raise FileNotFoundError(f"no ITSM workbook found in {search_dir} (tried {', '.join(_CANDIDATES)})")


def _sheet_prefix(label: str) -> str | None:
    m = re.match(r"\s*(\d+)(?!\d)", label or "")
    return m.group(1) if m else None


def _index_entries(wb) -> list[tuple[str, str, str]]:
    """[(sheet_label, prefix, entity_key)] from the Index sheet, in Index order."""
    if _INDEX_SHEET not in wb.sheetnames:
        raise ValueError(f"workbook has no '{_INDEX_SHEET}' sheet - cannot learn entity keys")
    ws = wb[_INDEX_SHEET]
    out: list[tuple[str, str, str]] = []
    header_seen = False
    for row in ws.iter_rows(values_only=True):
        a = _clean(row[0]) if len(row) > 0 else None
        b = _clean(row[1]) if len(row) > 1 else None
        if not header_seen:
            if isinstance(a, str) and a.lower() == "sheet":
                header_seen = True
            continue
        if not a or not b:
            continue
        prefix = _sheet_prefix(str(a))
        if prefix is None:
            continue
        key = re.sub(r"\.json$", "", str(b).strip(), flags=re.I)
        out.append((str(a), prefix, key))
    if not out:
        raise ValueError("Index sheet lists no entity keys")
    return out


def _find_sheet(wb, label: str, prefix: str):
    if label in wb.sheetnames:
        return wb[label]
    for ws in wb.worksheets:
        if ws.title != _INDEX_SHEET and _sheet_prefix(ws.title) == prefix:
            return ws
    return None


def _sheet_rows(ws) -> list[dict]:
    rows_iter = ws.iter_rows(values_only=True)
    headers: list[str] | None = None
    out: list[dict] = []
    for row in rows_iter:
        if headers is None:
            if not any(c is not None and str(c).strip() for c in row):
                continue
            headers = [snake(c) if c is not None and str(c).strip() else "" for c in row]
            continue
        vals = [_clean(c) for c in row]
        if not any(v is not None for v in vals):
            continue
        rec = {h: vals[i] for i, h in enumerate(headers) if h and i < len(vals)}
        out.append(rec)
    return out


def _bundle_version(wb, entities: dict, path: str) -> str:
    # An explicit version on the Read Me sheet wins.
    for rec in entities.get("meta", []) or []:
        if str(rec.get("item", "")).strip().lower() in ("version", "bundle version", "bundle_version"):
            v = rec.get("value")
            if v:
                return str(v)[:60]
    stem = os.path.splitext(os.path.basename(path))[0]
    m = re.search(r"_(v\d+)$", stem)
    token = m.group(1) if m else ("enriched" if stem.endswith("enriched") else "base")
    modified = getattr(wb.properties, "modified", None)
    if isinstance(modified, (_dt.datetime, _dt.date)):
        date = modified.date().isoformat() if isinstance(modified, _dt.datetime) else modified.isoformat()
    else:
        date = _dt.date.fromtimestamp(os.path.getmtime(path)).isoformat()
    return f"{token}-{date}"[:60]


# --- public API ---------------------------------------------------------------
def export_workbook(xlsx_path: str) -> dict:
    """Read the workbook and return the bundle dict (does not write a file)."""
    import openpyxl  # lazy: importer must not need it

    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=False)
    entities: dict[str, list[dict]] = {}
    for label, prefix, key in _index_entries(wb):
        ws = _find_sheet(wb, label, prefix)
        if ws is None:
            raise ValueError(f"Index lists sheet '{label}' ({key}) but the workbook has no such sheet")
        if key in entities:
            raise ValueError(f"duplicate entity key '{key}' in Index")
        entities[key] = _sheet_rows(ws)
    return {
        "bundle_version": _bundle_version(wb, entities, xlsx_path),
        "generated_at": _dt.datetime.now(_dt.timezone.utc).isoformat(),
        "source_filename": os.path.basename(xlsx_path),
        "checksum": compute_checksum(entities),
        "entities": entities,
    }


def write_bundle(bundle: dict, out_path: str) -> None:
    d = os.path.dirname(os.path.abspath(out_path))
    os.makedirs(d, exist_ok=True)
    with open(out_path, "w", encoding="utf-8", newline="\n") as fh:
        json.dump(bundle, fh, ensure_ascii=False, indent=2)
        fh.write("\n")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Export the ITSM workbook to a JSON bundle.")
    ap.add_argument("paths", nargs="+", metavar="PATH",
                    help="<xlsx> <out.json>  (or just <out.json> to auto-pick the newest workbook)")
    args = ap.parse_args(argv)
    if len(args.paths) == 1:
        xlsx, out = pick_workbook(None), args.paths[0]
    elif len(args.paths) == 2:
        xlsx, out = pick_workbook(args.paths[0]), args.paths[1]
    else:
        ap.error("expected <xlsx> <out.json>")
        return 2
    bundle = export_workbook(xlsx)
    write_bundle(bundle, out)
    counts = {k: len(v) for k, v in bundle["entities"].items()}
    print(f"bundle {bundle['bundle_version']} from {bundle['source_filename']} "
          f"-> {out}  checksum={bundle['checksum'][:12]}")
    for k, n in counts.items():
        print(f"  {k}: {n}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
