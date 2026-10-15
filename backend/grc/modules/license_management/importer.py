"""Basic CSV / XLSX import of entitlements (a purchase-order or licence-sheet
export). Headers are matched by synonym, every row goes through the same
validation as the register's add form, and re-importing the same sheet is
idempotent (identical rows are skipped). `dry_run=True` validates and reports
without writing."""
from __future__ import annotations

import csv
import io
import re
from typing import Any, Dict, List, Optional, Tuple

from sqlalchemy.orm import Session

from . import matching as M
from . import service as S
from .models import LicenseEntitlement

MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 60 * 1024 * 1024
MAX_ERRORS_RETURNED = 100

# canonical field -> accepted header spellings (normalised: lowercase, alphanumerics only),
# in priority order when a sheet carries several candidates.
_SYNONYMS: Dict[str, Tuple[str, ...]] = {
    "publisher": ("publisher", "vendor", "manufacturer", "softwarevendor", "softwarepublisher", "maker", "brand"),
    "product": ("product", "productname", "software", "softwarename", "application", "applicationname",
                "licensename", "sku", "skuname", "item", "title", "name", "license"),
    "edition": ("edition", "skuedition", "tier", "plan"),
    "license_type": ("licensetype", "licensemodel", "licensekind", "type", "model"),
    "license_metric": ("licensemetric", "metric", "licensedby", "countedby", "licensebasis", "unitofmeasure", "uom"),
    "total_quantity": ("totalquantity", "quantity", "qty", "seats", "licenses", "licences", "licensecount",
                       "numberoflicenses", "purchasedquantity", "entitled", "units", "count"),
    "unit_cost": ("unitcost", "unitprice", "priceperunit", "costperunit", "costperlicense", "costeach",
                  "price", "cost"),
    "currency": ("currency", "ccy", "curr"),
    "po_reference": ("poreference", "ponumber", "purchaseorder", "purchaseordernumber", "poref", "pono",
                     "ordernumber", "orderno", "po"),
    "contract_reference": ("contractreference", "contractnumber", "contractref", "agreementnumber",
                           "agreement", "enrollment", "contract"),
    "reseller": ("reseller", "supplier", "distributor", "partner", "purchasedfrom"),
    "purchase_date": ("purchasedate", "dateofpurchase", "orderdate", "invoicedate", "purchased", "boughton"),
    "start_date": ("startdate", "validfrom", "effectivedate", "termstart", "activationdate", "start"),
    "expiry_date": ("expirydate", "expirationdate", "expiry", "expiration", "expires", "enddate", "validuntil",
                    "validto", "renewaldate", "termend", "end"),
    "maintenance_expiry": ("maintenanceexpiry", "maintenanceexpiration", "maintenanceend", "maintenanceenddate",
                           "supportexpiry", "supportend", "supportenddate", "saexpiry"),
    "maintenance_included": ("maintenanceincluded", "hasmaintenance", "supportincluded", "maintenance"),
    "support_level": ("supportlevel", "supporttier", "supportplan"),
    "downgrade_rights": ("downgraderights", "downgrade"),
    "software_key": ("softwarekey", "matchkey"),
    "product_key": ("productkey", "licensekey", "activationkey", "serialnumber", "serial", "cdkey", "key"),
    "status": ("status", "state"),
    "notes": ("notes", "note", "comments", "comment", "remarks", "description"),
}

TEMPLATE_HEADERS = [
    "publisher", "product", "edition", "license_type", "license_metric", "total_quantity", "unit_cost",
    "currency", "po_reference", "contract_reference", "reseller", "purchase_date", "start_date",
    "expiry_date", "maintenance_expiry", "notes",
]
_TEMPLATE_ROWS = [
    ["Microsoft", "Microsoft 365 Apps for enterprise", "", "subscription", "per_user_named", "250", "180.00",
     "USD", "PO-2026-0412", "EA-118842", "Contoso Reseller", "2026-01-15", "2026-02-01", "2027-01-31", "",
     "Annual true-up in January"],
    ["Microsoft", "SQL Server 2022", "Standard", "perpetual", "per_core", "16", "3586.00",
     "USD", "PO-2025-0977", "", "Contoso Reseller", "2025-09-02", "2025-09-02", "", "2028-09-01",
     "Software Assurance until 2028"],
]


def template_csv() -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(TEMPLATE_HEADERS)
    w.writerows(_TEMPLATE_ROWS)
    return buf.getvalue()


def _norm_header(h: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def map_headers(headers: List[str]) -> Dict[int, str]:
    """column index -> canonical field. Each column feeds at most one field."""
    by_norm: Dict[str, int] = {}
    for i, h in enumerate(headers):
        n = _norm_header(h)
        if n and n not in by_norm:
            by_norm[n] = i
    mapping: Dict[int, str] = {}
    for field, syns in _SYNONYMS.items():
        for syn in syns:
            i = by_norm.get(syn)
            if i is not None and i not in mapping:
                mapping[i] = field
                break
    return mapping


def _read_csv(content: bytes) -> List[List[Any]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(r) for r in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(content: bytes) -> List[List[Any]]:
    import zipfile
    from openpyxl import load_workbook
    # An .xlsx is a zip: refuse one that would inflate far beyond its upload size.
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("The workbook is too large once unpacked.")
    except zipfile.BadZipFile:
        raise ValueError("Not a valid .xlsx workbook.")
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        rows: List[List[Any]] = []
        for r in wb.worksheets[0].iter_rows(values_only=True):
            rows.append(list(r))
            if len(rows) > MAX_ROWS + 200:          # header / blank rows allowance
                raise ValueError(f"Too many rows; the limit is {MAX_ROWS} per import.")
        return rows
    finally:
        wb.close()


def _blank(row: List[Any]) -> bool:
    return all(c is None or str(c).strip() == "" for c in row)


def parse_table(content: bytes, filename: str) -> Tuple[List[str], List[Tuple[int, List[Any]]]]:
    """(headers, [(sheet_row_number, cells)]). The header row is the first row
    with at least two non-empty cells."""
    if len(content) > MAX_BYTES:
        raise ValueError("File too large (max 5 MB).")
    if not content:
        raise ValueError("Empty file.")
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        rows = _read_xlsx(content)
    elif name.endswith(".xls"):
        raise ValueError("Legacy .xls is not supported - save the sheet as .xlsx or .csv.")
    else:
        rows = _read_csv(content)
    header_idx = next((i for i, r in enumerate(rows)
                       if sum(1 for c in r if c is not None and str(c).strip() != "") >= 2), None)
    if header_idx is None:
        raise ValueError("No header row found.")
    headers = [str(c).strip() if c is not None else "" for c in rows[header_idx]]
    body = [(i + 1, r) for i, r in enumerate(rows) if i > header_idx and not _blank(r)]
    return headers, body


def _natural_key(f: Dict[str, Any]) -> tuple:
    exp = f.get("expiry_date")
    return (
        M.clean(f.get("publisher")).lower(), M.clean(f.get("product")).lower(),
        M.clean(f.get("edition")).lower(), f.get("license_metric"),
        M.clean(f.get("po_reference")).lower(), int(f.get("total_quantity") or 0),
        exp.date().isoformat() if exp else "",
    )


def import_entitlements(db: Session, tenant_id: int, content: bytes, filename: str,
                        user: Optional[str] = None, dry_run: bool = False) -> Dict[str, Any]:
    S.ensure_tables(db)
    headers, body = parse_table(content, filename)
    if len(body) > MAX_ROWS:
        raise ValueError(f"Too many rows ({len(body)}); the limit is {MAX_ROWS} per import.")
    mapping = map_headers(headers)
    fields = set(mapping.values())
    missing = [f for f in ("publisher", "product") if f not in fields]
    if missing:
        raise ValueError(
            "Could not find a " + " and a ".join(
                "Publisher (or Vendor) column" if m == "publisher" else "Product (or Software) column"
                for m in missing) + " in the header row. Download the template for the expected columns.")

    existing = {
        _natural_key({
            "publisher": l.publisher, "product": l.product, "edition": l.edition,
            "license_metric": l.license_metric, "po_reference": l.po_reference,
            "total_quantity": l.total_quantity, "expiry_date": l.expiry_date,
        })
        for l in db.query(LicenseEntitlement).filter(LicenseEntitlement.tenant_id == tenant_id)
    }

    errors: List[Dict[str, Any]] = []
    error_count = 0
    skipped = 0
    to_add: List[LicenseEntitlement] = []
    for rownum, cells in body:
        data: Dict[str, Any] = {}
        for i, field in mapping.items():
            if i < len(cells) and cells[i] is not None and str(cells[i]).strip() != "":
                data[field] = cells[i]
        try:
            f = S.clean_entitlement(data, partial=False)
            S._check_dates(f.get("start_date"), f.get("expiry_date"))
            plain = f.pop("_product_key", "")
            nk = _natural_key(f)
            if nk in existing:
                skipped += 1
                continue
            existing.add(nk)
            row = LicenseEntitlement(tenant_id=tenant_id, created_by=user or None, **f)
            row.status = f.get("status", "active")
            if not dry_run:
                row.product_key = S._encrypt_key(plain)
            to_add.append(row)
        except S.EncryptionUnavailable:
            raise
        except ValueError as exc:
            error_count += 1
            if len(errors) < MAX_ERRORS_RETURNED:
                errors.append({"row": rownum, "error": str(exc)})

    created_ids: List[int] = []
    if not dry_run and to_add:
        db.add_all(to_add)
        db.commit()
        created_ids = [r.id for r in to_add]
        S._rebuild_quietly(db, tenant_id)
    return {
        "dry_run": dry_run,
        "total": len(body),
        "created": len(to_add) if not dry_run else 0,
        "would_create": len(to_add) if dry_run else None,
        "skipped": skipped,
        "error_count": error_count,
        "errors": errors,
        "ids": created_ids,
        "columns": {headers[i]: f for i, f in mapping.items()},
        "unmapped": [h for i, h in enumerate(headers) if h and i not in mapping],
    }
