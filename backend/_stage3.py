import json, os
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

HERE = os.path.dirname(__file__)
SEED = os.path.join(HERE, 'grc', 'seed_data')

# Pre-2 base (controls under categories: full detail + evidence + artifacts) + Stage-2 map (category -> master domain)
pre2 = json.load(open(os.path.join(SEED, 'pre2_controls_by_framework_category.json'), encoding='utf-8'))
s2 = json.load(open(os.path.join(SEED, 'stage2_master_categories.json'), encoding='utf-8'))
cat2master = {}
for m in s2['master_categories']:
    for it in m['merged_from']:
        cat2master[it['name']] = m['name']
order = [m['name'] for m in s2['master_categories']]

by_master = defaultdict(list)
total = unmapped = tot_ev = tot_art = 0
for item in pre2['items']:
    master = cat2master.get(item['category'])
    if master is None:
        unmapped += 1
        master = 'UNMAPPED'
        if master not in order:
            order.append(master)
    for c in item['controls']:
        total += 1
        ev = c.get('evidence') or []
        art = c.get('artifacts') or []
        tot_ev += len(ev); tot_art += len(art)
        by_master[master].append({
            'framework': item['framework'], 'framework_version': item['framework_version'],
            'control_id': c['control_id'], 'reference': c['reference'],
            'section_number': c['section_number'], 'parent_section': c['parent_section'],
            'original_category': item['category'], 'title': c['title'],
            'evidence': ev, 'artifacts': art,
        })
placed = sum(len(v) for v in by_master.values())
print('total controls :', total, ' placed :', placed, ' unmapped categories :', unmapped)
print('evidence carried:', tot_ev, ' artifacts carried:', tot_art)

# JSON (system)
out = {'stage': 3,
       'desc': 'controls under the 20 master domains, FULL source detail + evidence + artifacts (domains merged; controls NOT yet normalized)',
       'total_controls': total, 'placed': placed, 'evidence_carried': tot_ev, 'artifacts_carried': tot_art,
       'domains': [{'domain': n, 'control_count': len(by_master.get(n, [])), 'controls': by_master.get(n, [])} for n in order],
       'framework_level_artifacts_by_framework': pre2.get('framework_level_artifacts_by_framework', {})}
jpath = os.path.join(SEED, 'stage3_controls_by_domain.json')
json.dump(out, open(jpath, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

# Excel (human)
wb = openpyxl.Workbook(); hd = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='2E7D6B'); thin = Side(style='thin', color='DDDDDD'); band = PatternFill('solid', fgColor='EAF3F0')
# Sheet 1: summary per domain
sm = wb.active; sm.title = 'Summary'
for j, h in enumerate(['Master Domain', '# Controls', '# Evidences', '# Artifacts'], 1):
    cc = sm.cell(1, j, h); cc.font = hd; cc.fill = fill
r = 2
for n in order:
    recs = by_master.get(n, [])
    ne = sum(len(x['evidence']) for x in recs); na = sum(len(x['artifacts']) for x in recs)
    sm.cell(r, 1, n); sm.cell(r, 2, len(recs)); sm.cell(r, 3, ne); sm.cell(r, 4, na)
    for col in (1, 2, 3, 4): sm.cell(r, col).border = Border(bottom=thin)
    r += 1
sm.cell(r, 1, 'TOTAL').font = Font(bold=True); sm.cell(r, 2, placed).font = Font(bold=True)
sm.cell(r, 3, tot_ev).font = Font(bold=True); sm.cell(r, 4, tot_art).font = Font(bold=True)
sm.freeze_panes = 'A2'
for col, w in zip('ABCD', [42, 12, 12, 12]): sm.column_dimensions[col].width = w
# Sheet 2: every control under its domain, full detail + evidence + artifacts
de = wb.create_sheet('Controls by Domain (full)')
heads = ['Master Domain', 'Framework', 'Version', 'Control ID', 'Reference / Clause', 'Section #', 'Parent Section',
         'Original Category', 'Control Title', '# Ev', 'Recommended Evidences (per control)', '# Art', 'Artifacts (per control)']
for j, h in enumerate(heads, 1):
    cc = de.cell(1, j, h); cc.font = hd; cc.fill = fill; cc.alignment = Alignment(vertical='center', wrap_text=True)
r = 2; flip = False; prev = None
for n in order:
    for rec in by_master.get(n, []):
        if n != prev: flip = not flip; prev = n
        evlist = rec['evidence']; artlist = rec['artifacts']
        evtext = '\n'.join('• ' + e['name'] for e in evlist if e.get('name'))
        attext = '\n'.join('• ' + (a.get('name') or '') + ' [' + str(a.get('type') or '') + ']' for a in artlist)
        vals = [n, rec['framework'], rec['framework_version'], rec['control_id'], rec['reference'], rec['section_number'],
                rec['parent_section'], rec['original_category'], rec['title'], len(evlist), evtext, len(artlist), attext]
        for j, v in enumerate(vals, 1):
            cell = de.cell(r, j, v)
            if j in (11, 13): cell.alignment = Alignment(wrap_text=True, vertical='top')
        if flip:
            for col in range(1, 14): de.cell(r, col).fill = band
        for col in range(1, 14): de.cell(r, col).border = Border(bottom=thin)
        r += 1
de.freeze_panes = 'A2'
for col, w in zip('ABCDEFGHIJKLM', [28, 26, 8, 12, 24, 9, 11, 24, 44, 6, 60, 6, 55]):
    de.column_dimensions[col].width = w
# Sheet 3: framework-level artifacts (not pinned to one control)
s3 = wb.create_sheet('Framework-level Artifacts')
for j, h in enumerate(['Framework', 'Artifact Name', 'Type', 'Format', 'control_ref'], 1):
    cc = s3.cell(1, j, h); cc.font = hd; cc.fill = fill
r = 2
for fw, lst in pre2.get('framework_level_artifacts_by_framework', {}).items():
    for a in lst:
        s3.cell(r, 1, fw); s3.cell(r, 2, a.get('name')); s3.cell(r, 3, a.get('type')); s3.cell(r, 4, a.get('format')); s3.cell(r, 5, a.get('control_ref'))
        for col in range(1, 6): s3.cell(r, col).border = Border(bottom=thin)
        r += 1
s3.freeze_panes = 'A2'
for col, w in zip('ABCDE', [34, 50, 16, 14, 24]): s3.column_dimensions[col].width = w

xpath = os.path.expanduser('~/OneDrive/Desktop/STAGE3_controls_by_domain.xlsx')
try: wb.save(xpath)
except PermissionError: xpath = os.path.expanduser('~/OneDrive/Desktop/STAGE3_controls_by_domain_v2.xlsx'); wb.save(xpath)
print('JSON  (system):', jpath)
print('EXCEL (you)   :', xpath)
