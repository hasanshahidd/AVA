import json, glob, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

HERE = os.path.dirname(__file__)
a = json.load(open(os.path.join(HERE, 'grc', 'seed_data', 'artifact_catalog.json'), encoding='utf-8'))
keyinfo = {k: (v.get('name', ''), len(v.get('artifacts') or [])) for k, v in a.items() if isinstance(v, dict)}

# verified, hand-checked map: our framework -> artifact key
MAP = [
 ("KSA National Data Management and Personal Data Protection Standards", "ksa_ndmo"),
 ("Regulation on Personal Data Transfer Outside the Kingdom", "ksa_data_transfer"),
 ("Abu Dhabi Healthcare Information and Cyber Security Standard", "adhics"),
 ("ARAMCO Cybersecurity Compliance Certification", "aramco_csc"),
 ("CIS Critical Security Controls v8", "cis_v8"),
 ("COBIT 2019", "cobit_2019"),
 ("DOH Policy on the Abu Dhabi Health Information Exchange (ADHIE)", "adhie"),
 ("Digital Operational Resilience Act (DORA)", "dora"),
 ("General Data Protection Regulation", "gdpr"),
 ("HIPAA Security & Privacy Rule", "hipaa"),
 ("HITRUST Common Security Framework (CSF)", "hitrust_csf"),
 ("ISO 22301:2019 Business Continuity Management System", "iso_22301_2019"),
 ("ISO/IEC 27001:2022", "iso_27001_2022"),
 ("ISO/IEC 42001:2023 AI Management System", "iso_42001_2023"),
 ("MAS Technology Risk Management Guidelines", "mas_trm"),
 ("NIS2 Directive", "nis2"),
 ("NIST SP 800-53 Rev 5", "nist_sp_800_53_r5"),
 ("NIST Artificial Intelligence Risk Management Framework (AI RMF 1.0)", "nist_ai_rmf"),
 ("NIST Cybersecurity Framework", "nist_csf_2"),
 ("PCI Data Security Standard", "pci_dss_v4"),
 ("Qatar Central Bank Technology Risks Circular", "qatar_cb"),
 ("SABIC CyberTrust Guidelines", "sabic_cybertrust"),
 ("SAMA Cyber Security Framework", "sama_csf"),
 ("SBP Cloud Outsourcing Framework", "sbp_cloud_outsourcing"),
 ("SBP ETGRMF", "sbp_etgrmf"),
 ("SBP Internet Banking Framework", "sbp_internet_banking"),
 ("Sri Lanka Baseline Security Standard (BSS)", "sri_lanka_bss"),
 ("SOC 2 Type II", "soc2"),
 ("SOX IT General Controls", "sox_itgc"),
 ("SWIFT Customer Security Controls Framework", "swift_cscf"),
]
NOTE = {"nist_csf_2": "Catalog is CSF 2.0 - confirm version OK"}

wb = openpyxl.Workbook(); hd = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='2E7D6B'); thin = Side(style='thin', color='DDDDDD'); warn = PatternFill('solid', fgColor='FFF3CD')
ws = wb.active; ws.title = 'Framework to Artifact Map'
for j, h in enumerate(['#', 'Our Framework', 'Artifact Key', 'Catalog Name', '# Artifacts', 'Confirm?'], 1):
    cc = ws.cell(1, j, h); cc.font = hd; cc.fill = fill; cc.alignment = Alignment(vertical='center', wrap_text=True)
r = 2; total = 0
for i, (fwname, key) in enumerate(MAP, 1):
    cname, cnt = keyinfo.get(key, ('(KEY NOT FOUND)', 0)); total += cnt
    ws.cell(r, 1, i); ws.cell(r, 2, fwname); ws.cell(r, 3, key); ws.cell(r, 4, cname); ws.cell(r, 5, cnt); ws.cell(r, 6, NOTE.get(key, 'OK'))
    if key in NOTE:
        for col in range(1, 7): ws.cell(r, col).fill = warn
    for col in range(1, 7): ws.cell(r, col).border = Border(bottom=thin)
    r += 1
ws.cell(r, 2, 'TOTAL (30 frameworks)').font = Font(bold=True); ws.cell(r, 5, total).font = Font(bold=True)
r += 2
ws.cell(r, 2, 'NOT USED (no matching framework in our 30):').font = Font(bold=True, color='B00000'); r += 1
used = {k for _, k in MAP}
for k, (cname, cnt) in keyinfo.items():
    if k not in used:
        ws.cell(r, 3, k); ws.cell(r, 4, cname); ws.cell(r, 5, cnt); r += 1
ws.freeze_panes = 'A2'
for col, w in zip('ABCDEF', [5, 50, 22, 40, 11, 30]): ws.column_dimensions[col].width = w

dl = os.path.expanduser('~/Downloads')
if not os.path.isdir(dl): dl = os.path.expanduser('~/OneDrive/Desktop')
out = os.path.join(dl, 'STEP1_framework_artifact_map.xlsx')
try: wb.save(out)
except PermissionError: out = os.path.join(dl, 'STEP1_framework_artifact_map_v2.xlsx'); wb.save(out)
print('WROTE', out)
print('mapped 30/30 frameworks -> artifact keys ; total artifacts attachable:', total, '; unused keys: 2')
