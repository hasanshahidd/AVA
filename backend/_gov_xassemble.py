import json, os, re
from collections import OrderedDict, Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

HERE = os.path.dirname(__file__); SEED = os.path.join(HERE, 'grc', 'seed_data')
gov = [d for d in json.load(open(os.path.join(SEED, 'stage3_complete.json'), encoding='utf-8'))['domains'] if d['domain'].startswith('Governance')][0]['controls']
sets = json.load(open(os.path.join(HERE, '_gov_xsets.json'), encoding='utf-8'))['sets']
ev_norm = {}
for fn in ['_gov_xev_00_20.json', '_gov_xev_21_40.json']:
    for s in json.load(open(os.path.join(HERE, fn), encoding='utf-8'))['sets']:
        ev_norm[s['set_id']] = s['normalized_evidence']

def nk(s): return re.sub(r'\s+', ' ', re.sub(r'[^a-z0-9 ]', ' ', (s or '').lower())).strip()
def fwshort(c): return c['framework'].split('(')[0].strip()
def own_ev(i):
    g = OrderedDict()
    for e in (gov[i].get('evidence') or []):
        k = nk(e['name'])
        if k and k not in g: g[k] = {'name': e['name'], 'absorbs': [e['name']], 'sources': [fwshort(gov[i]) + ' ' + gov[i]['control_id']]}
    return list(g.values())
def arts(idxs):
    g = OrderedDict()
    for i in idxs:
        for a in (gov[i].get('artifacts') or []):
            k = (nk(a.get('name')), (a.get('type') or '').lower())
            if not k[0]: continue
            if k not in g: g[k] = {'name': a.get('name'), 'type': a.get('type'), 'absorbs': 0, 'sources': set()}
            g[k]['absorbs'] += 1; g[k]['sources'].add(fwshort(gov[i]) + ' ' + gov[i]['control_id'])
    return [{'name': v['name'], 'type': v['type'], 'absorbs': v['absorbs'], 'sources': sorted(v['sources'])} for v in g.values()]

out_sets = []
for s in sets:
    idxs = s['members']
    members = [{'framework': fwshort(gov[i]), 'control_id': gov[i]['control_id'], 'original_title': gov[i]['title'], 'reference': gov[i].get('reference', '')} for i in idxs]
    ne = ev_norm.get(s['set_id']) if len(idxs) > 1 else own_ev(idxs[0])
    out_sets.append({'set_id': s['set_id'], 'normalized_title': s['normalized_title'], 'member_count': len(members),
                     'frameworks': sorted(set(m['framework'] for m in members)), 'members': members,
                     'normalized_evidence': ne or [], 'normalized_artifacts': arts(idxs)})

# ---- INVARIANTS ----
allm = [i for s in sets for i in s['members']]
print('VERIFY coverage 420:', 'OK' if Counter(allm) == Counter(range(420)) else 'FAIL')
viol = [s['set_id'] for s in out_sets if len(s['frameworks']) != s['member_count']]
print('VERIFY no same-framework in a set:', 'OK (0)' if not viol else f'FAIL {viol[:8]}')
tmis = sum(1 for s, src in zip(out_sets, sets) for m, i in zip(s['members'], src['members']) if m['original_title'] != gov[i]['title'])
print('VERIFY titles unchanged:', 'OK' if tmis == 0 else f'{tmis} mismatch')

multi = [s for s in out_sets if s['member_count'] > 1]; single = [s for s in out_sets if s['member_count'] == 1]
import statistics
evc = [len(s['normalized_evidence']) for s in multi]
raw_ev = sum(len(c.get('evidence') or []) for c in gov); norm_ev = sum(len(s['normalized_evidence']) for s in out_sets)
raw_art = sum(len(c.get('artifacts') or []) for c in gov); norm_art = sum(len(s['normalized_artifacts']) for s in out_sets)
print('sets:', len(out_sets), '(', len(multi), 'normalized +', len(single), 'standalone )')
print('normalized-set evidence: median', int(statistics.median(evc)), 'max', max(evc), '| >8:', sum(1 for x in evc if x > 8))
print('frameworks/set: max', max(len(s['frameworks']) for s in multi))
print('evidence %d->%d | artifacts %d->%d' % (raw_ev, norm_ev, raw_art, norm_art))

out_sets.sort(key=lambda s: (-s['member_count'], s['set_id']))
out = {'stage': 4, 'version': '4-xframework', 'domain': 'Governance, Leadership & Policy', 'controls_in': 420, 'sets_out': len(out_sets),
       'normalized_sets': len(multi), 'standalone': len(single), 'frameworks': sorted(set(f for s in out_sets for f in s['frameworks'])),
       'evidence_raw': raw_ev, 'evidence_normalized': norm_ev, 'artifacts_raw': raw_art, 'artifacts_normalized': norm_art, 'sets': out_sets}
json.dump(out, open(os.path.join(SEED, 'stage4_sets_governance.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

wb = openpyxl.Workbook(); hd = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='2E7D6B'); thin = Side(style='thin', color='DDDDDD')
sm = wb.active; sm.title = 'Sets summary'
for j, h in enumerate(['Set ID', 'Normalized Heading', '# Controls=Frameworks', '# Evidence', '# Artifacts', 'Type'], 1):
    cc = sm.cell(1, j, h); cc.font = hd; cc.fill = fill; cc.alignment = Alignment(wrap_text=True, vertical='center')
r = 2
for s in out_sets:
    sm.cell(r, 1, s['set_id']); sm.cell(r, 2, s['normalized_title']); sm.cell(r, 3, s['member_count'])
    sm.cell(r, 4, len(s['normalized_evidence'])); sm.cell(r, 5, len(s['normalized_artifacts'])); sm.cell(r, 6, 'standalone' if s['member_count'] == 1 else 'normalized set')
    for col in range(1, 7): sm.cell(r, col).border = Border(bottom=thin)
    r += 1
sm.freeze_panes = 'A2'
for col, w in zip('ABCDEF', [12, 56, 16, 10, 10, 14]): sm.column_dimensions[col].width = w
xp = os.path.expanduser('~/OneDrive/Desktop/GOVERNANCE_normalized.xlsx')
try: wb.save(xp)
except PermissionError: xp = os.path.expanduser('~/OneDrive/Desktop/GOVERNANCE_normalized_v2.xlsx'); wb.save(xp)
print('JSON :', os.path.join(SEED, 'stage4_sets_governance.json')); print('EXCEL:', xp)
