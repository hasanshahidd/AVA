"""Software license management — recognition, reconciliation, summary, CRUD,
import and wiring.

DB tests run against an in-memory SQLite DB holding only the three tables the
engine touches (FKs to uncreated tables are inert on SQLite). The PostgreSQL-only
parts (advisory lock, partial-index DDL) are exercised by the dialect guard, not
by SQLite.
"""
from datetime import datetime, timedelta

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from grc.models import Base, ITAsset
from grc.modules.license_management import importer, matching as M, service as S
from grc.modules.license_management.models import LicenseAssignment, LicenseEntitlement

TID = 1
NOW = datetime(2026, 10, 15, 12, 0, 0)


# ── fixtures ────────────────────────────────────────────────────────────────

@pytest.fixture()
def db(monkeypatch):
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine, tables=[ITAsset.__table__, LicenseEntitlement.__table__,
                                             LicenseAssignment.__table__])
    S._ENSURED.clear()
    S._LAST_REBUILD.clear()
    S._LAST_STATS.clear()
    monkeypatch.setattr(S, "REBUILD_TTL_SECONDS", 0)      # every read re-derives
    monkeypatch.setattr(S, "_now", lambda: NOW)
    s = sessionmaker(bind=engine)()
    yield s
    s.close()


def sw(name, version="1.0", publisher="", source="registry", key=None):
    e = {"name": name, "version": version, "source": source}
    if publisher:
        e["publisher"] = publisher
    if key:
        e["software_key"] = key
    return e


def asset(db, name, software, **kw):
    a = ITAsset(tenant_id=TID, name=name, asset_type="infrastructure",
                detected_software_json=software, last_seen_at=NOW - timedelta(days=1), **kw)
    db.add(a)
    db.commit()
    return a


def ent(db, **kw):
    d = dict(publisher="Microsoft", product="Office", license_type="perpetual",
             license_metric="per_device", total_quantity=10)
    d.update(kw)
    return S.create_entitlement(db, TID, d, user="t@x", rebuild=False)


def pos(db, product):
    out = S.compute_license_positions(db, TID)
    rows = [p for p in out["positions"] if p["product"] == product]
    assert len(rows) == 1, [p["product"] for p in out["positions"]]
    return rows[0]


OFFICE = "Microsoft Office Professional Plus 2019"


# ── recognition (pure) ──────────────────────────────────────────────────────

def test_name_tokens_strip_noise_but_keep_identity():
    assert M.name_tokens("7-Zip 23.01 (x64)") == ["7", "zip"]
    assert M.name_tokens("Microsoft SQL Server 2022 (64-bit)") == ["microsoft", "sql", "server", "2022"]
    assert M.name_tokens("Microsoft Office Professional Plus 2019 - en-us") == \
        ["microsoft", "office", "professional", "plus", "2019"]
    assert M.name_tokens("Notepad++ (64-bit x64)") == ["notepad++"]


def test_locale_variants_share_one_name_key():
    assert M.name_key("Microsoft Office Professional Plus 2019 - en-us") == \
        M.name_key("Microsoft Office Professional Plus 2019 - de-de")


def _tier(publisher, product, edition, install_name, install_pub="", keys=(), install_key=None):
    d = M.Descriptor(publisher, product, edition, keys)
    inst = M.read_install({"name": install_name, "publisher": install_pub, "source": "registry",
                           **({"software_key": install_key} if install_key else {})})
    return d.tier(inst)


def test_product_matches_registry_name_with_publisher_prefix_and_arch_noise():
    assert _tier("Microsoft", "SQL Server 2022", None, "Microsoft SQL Server 2022 (64-bit)", "Microsoft Corporation") is not None
    assert _tier("Adobe", "Acrobat Reader", None, "Adobe Acrobat Reader DC", "Adobe Inc.") is not None


def test_visual_studio_never_swallows_visual_studio_code():
    assert _tier("Microsoft", "Visual Studio", None, "Microsoft Visual Studio Code", "Microsoft Corporation") is None
    assert _tier("Microsoft", "Visual Studio", "Professional", "Microsoft Visual Studio Professional 2022") is not None


def test_different_known_publishers_never_match():
    assert _tier("Adobe", "Reader", None, "Reader", "Foxit Software") is None


def test_edition_conflict_blocks_but_unstated_edition_matches():
    assert _tier("Microsoft", "SQL Server 2022", "Enterprise", "Microsoft SQL Server 2022 Standard", "Microsoft") is None
    assert _tier("Microsoft", "SQL Server 2022", "Enterprise", "Microsoft SQL Server 2022 (64-bit)", "Microsoft") is not None


def test_pro_and_professional_are_the_same_edition():
    assert _tier("Microsoft", "Visual Studio", "Pro", "Visual Studio Professional 2022", "Microsoft") is not None


def test_different_year_is_a_different_product():
    assert _tier("Microsoft", "Windows Server 2019", None, "Windows Server 2022", "Microsoft") is None


def test_explicit_key_matches_exactly_and_disables_name_matching():
    d = M.Descriptor("Microsoft", "Whatever", None, ["mssql-2022"])
    assert d.fuzzy is False
    assert d.tier(M.read_install({"name": "Microsoft SQL Server 2022 (64-bit)", "source": "registry",
                                  "software_key": "mssql-2022"})) == 1
    assert d.tier(M.read_install({"name": "Whatever", "source": "registry", "software_key": "other"})) is None


def test_prefix_key_spec():
    d = M.Descriptor("Microsoft", "M365", None, ["microsoft-365-apps*"])
    assert d.tier(M.read_install({"name": "Microsoft 365 Apps for enterprise - en-us", "source": "registry"})) == 2


def test_read_install_skips_roles_processes_and_noise():
    assert M.read_install({"name": "Web-Server", "source": "windows_role"}) is None
    assert M.read_install({"name": "postgres", "source": "listening_process"}) is None
    assert M.read_install({"name": "Microsoft Visual C++ 2015 Redistributable (x64)", "source": "registry"}) is None
    assert M.read_install("not a dict") is None


def test_wmi_vendor_in_source_is_recovered_as_publisher():
    i = M.read_install({"name": "Some App", "source": "Contoso Ltd"})
    assert i["publisher"] == "Contoso Ltd"
    assert M.read_install({"name": "Some App", "source": "registry"})["publisher"] == ""


# ── engine: consumption + position ──────────────────────────────────────────

def test_per_device_counts_one_per_asset_and_position_goes_negative(db):
    ent(db, total_quantity=2)
    for i in range(3):
        asset(db, f"pc{i}", [sw(OFFICE, publisher="Microsoft Corporation")])
    p = pos(db, "Office")
    assert (p["entitled"], p["consumed"], p["position"]) == (2, 3, -1)
    assert p["status"] == "over_deployed"
    assert {a["quantity"] for a in p["consuming_assets"]} == {1}
    assert p["consuming_assets_total"] == 3


def test_exposure_is_shortfall_times_unit_cost_and_null_when_unpriced(db):
    ent(db, total_quantity=1, unit_cost=150.0)
    for i in range(3):
        asset(db, f"pc{i}", [sw(OFFICE, publisher="Microsoft")])
    p = pos(db, "Office")
    assert p["position"] == -2 and p["exposure_cost"] == 300.0

    db.query(LicenseEntitlement).update({LicenseEntitlement.unit_cost: None})
    db.commit()
    p = pos(db, "Office")
    assert p["exposure_cost"] is None and p["unit_cost"] is None   # unknown price is not zero


def test_compliant_when_entitled_covers_installs(db):
    ent(db, total_quantity=3)
    for i in range(3):
        asset(db, f"pc{i}", [sw(OFFICE, publisher="Microsoft")])
    p = pos(db, "Office")
    assert p["position"] == 0 and p["status"] == "compliant" and p["exposure_cost"] == 0.0


def test_surplus_beyond_buffer_is_over_licensed_and_reclaimable(db):
    ent(db, total_quantity=20, unit_cost=100.0)
    asset(db, "pc0", [sw(OFFICE, publisher="Microsoft")])
    p = pos(db, "Office")
    assert p["position"] == 19 and p["status"] == "over_licensed"
    assert p["reclaimable_count"] == 19 - 2          # 10% buffer of 20 = 2 seats kept
    s = S.compute_summary(db, TID)
    assert s["reclaimable_count"] == 17 and s["reclaimable_savings"] == 1700.0


def test_small_surplus_inside_buffer_stays_compliant(db):
    ent(db, total_quantity=10)
    for i in range(9):
        asset(db, f"pc{i}", [sw(OFFICE, publisher="Microsoft")])
    p = pos(db, "Office")
    assert p["position"] == 1 and p["status"] == "compliant" and p["reclaimable_count"] == 0


def test_expiring_soon_when_nearest_expiry_within_90_days(db):
    ent(db, total_quantity=1, expiry_date=NOW + timedelta(days=30))
    asset(db, "pc0", [sw(OFFICE, publisher="Microsoft")])
    p = pos(db, "Office")
    assert p["status"] == "expiring_soon" and p["nearest_expiry"] == (NOW + timedelta(days=30)).isoformat()


def test_over_deployed_outranks_expiring_soon(db):
    ent(db, total_quantity=1, expiry_date=NOW + timedelta(days=10))
    for i in range(2):
        asset(db, f"pc{i}", [sw(OFFICE, publisher="Microsoft")])
    assert pos(db, "Office")["status"] == "over_deployed"


def test_unlicensed_install_has_no_entitlement_and_is_listed(db):
    ent(db, product="Something Else", publisher="Contoso")
    asset(db, "pc0", [sw("Acme Draw 4", publisher="Acme Corp")])
    asset(db, "pc1", [sw("Acme Draw 4", publisher="Acme Corp")])
    out = S.compute_license_positions(db, TID)
    assert len(out["unlicensed"]) == 1
    u = out["unlicensed"][0]
    assert (u["name"], u["install_count"], u["publisher"]) == ("Acme Draw 4", 2, "Acme Corp")
    assert [a["name"] for a in u["assets"]] == ["pc0", "pc1"] and u["reason"] == "no_entitlement"
    assert S.compute_summary(db, TID)["unlicensed_count"] == 2
    rows = db.query(LicenseAssignment).filter(LicenseAssignment.entitlement_id.is_(None)).all()
    assert len(rows) == 2


def test_version_variants_of_one_unlicensed_product_collapse(db):
    asset(db, "pc0", [sw("7-Zip 23.01 (x64)", publisher="Igor Pavlov", key="7-zip-23-01")])
    asset(db, "pc1", [sw("7-Zip 22.01", publisher="Igor Pavlov", key="7-zip-22-01")])
    out = S.compute_license_positions(db, TID)["unlicensed"]
    assert len(out) == 1 and out[0]["install_count"] == 2


def test_locale_duplicates_on_one_asset_count_once(db):
    ent(db, total_quantity=5)
    asset(db, "pc0", [sw(OFFICE + " - en-us", publisher="Microsoft"),
                      sw(OFFICE + " - de-de", publisher="Microsoft")])
    assert pos(db, "Office")["consumed"] == 1


def test_roles_processes_and_linux_packages_are_not_flagged_unlicensed(db):
    asset(db, "srv", [
        sw("Web-Server", source="windows_role"),
        sw("postgres", source="listening_process"),
        sw("libc6", source="dpkg"),
        sw("bash", source="rpm"),
        sw("Microsoft Visual C++ 2019 Redistributable (x64)"),
    ])
    assert S.compute_license_positions(db, TID)["unlicensed"] == []


def test_linux_package_can_still_match_an_entitlement(db):
    ent(db, publisher="Splunk", product="splunk", total_quantity=1, software_key="splunk")
    asset(db, "lnx", [sw("splunk", version="9.1", source="rpm", key="splunk")])
    p = pos(db, "splunk")
    assert p["consumed"] == 1 and p["position"] == 0


def test_expired_entitlement_does_not_license_installs(db):
    ent(db, total_quantity=5, expiry_date=NOW - timedelta(days=30))
    asset(db, "pc0", [sw(OFFICE, publisher="Microsoft")])
    out = S.compute_license_positions(db, TID)
    assert out["unlicensed"][0]["reason"] == "entitlement_expired"
    assert pos(db, "Office")["status"] == "expired"
    s = S.compute_summary(db, TID)
    assert s["tracked_products"] == 0 and s["unlicensed_count"] == 1 and s["expired_lines"] == 1
    # the sweep also flipped the stored status so the report dataset agrees
    assert db.query(LicenseEntitlement).one().status == "expired"


def test_date_only_expiry_is_valid_through_that_day(db):
    ent(db, total_quantity=1, expiry_date=datetime(2026, 10, 15))     # NOW is 12:00 on the 15th
    asset(db, "pc0", [sw(OFFICE, publisher="Microsoft")])
    assert pos(db, "Office")["consumed"] == 1


def test_not_yet_started_line_is_pending_not_licensing(db):
    ent(db, total_quantity=5, start_date=NOW + timedelta(days=10))
    asset(db, "pc0", [sw(OFFICE, publisher="Microsoft")])
    assert pos(db, "Office")["status"] == "pending"
    assert S.compute_license_positions(db, TID)["unlicensed"][0]["install_count"] == 1


def test_multiple_lines_of_one_product_pool(db):
    ent(db, total_quantity=2)
    ent(db, total_quantity=3, po_reference="PO-2")
    for i in range(6):
        asset(db, f"pc{i}", [sw(OFFICE, publisher="Microsoft")])
    p = pos(db, "Office")
    assert (p["entitled"], p["consumed"], p["position"]) == (5, 6, -1)
    assert len(p["entitlement_ids"]) == 2
    lines = S.list_entitlements(db, TID)["entitlements"]
    assert [(l["consumed"], l["position"], l["status"]) for l in lines] == [(6, -1, "over_deployed")] * 2


def test_two_editions_of_one_key_fill_in_order_and_spill_over(db):
    ent(db, publisher="Microsoft", product="SQL Server 2022", edition="Standard", total_quantity=1)
    ent(db, publisher="Microsoft", product="SQL Server 2022", edition="Enterprise", total_quantity=1)
    for i in range(2):
        asset(db, f"db{i}", [sw("Microsoft SQL Server 2022 (64-bit)", publisher="Microsoft Corporation")])
    out = S.compute_license_positions(db, TID)["positions"]
    assert sorted((p["edition"], p["consumed"], p["position"]) for p in out) == \
        [("Enterprise", 1, 0), ("Standard", 1, 0)]


def test_per_core_uses_physical_cores_then_logical_then_one(db):
    ent(db, product="SQL Server 2022", license_metric="per_core", total_quantity=8)
    asset(db, "phys", [sw("Microsoft SQL Server 2022 (64-bit)", publisher="Microsoft")], cpu_cores=16,
          platform_properties={"cpu": {"status": "discovered",
                                       "data": {"physical_cores": 8, "logical_processors": 16, "sockets": 1}}})
    asset(db, "logical", [sw("Microsoft SQL Server 2022 (64-bit)", publisher="Microsoft")], cpu_cores=4)
    asset(db, "unknown", [sw("Microsoft SQL Server 2022 (64-bit)", publisher="Microsoft")])
    p = pos(db, "SQL Server 2022")
    qty = {a["name"]: a["quantity"] for a in p["consuming_assets"]}
    assert qty == {"phys": 8, "logical": 4, "unknown": 1}
    assert p["consumed"] == 13 and p["position"] == -5
    note = db.query(LicenseAssignment).join(ITAsset, ITAsset.id == LicenseAssignment.asset_id) \
        .filter(ITAsset.name == "unknown").one().notes
    assert "not collected" in note            # a guessed count is annotated, never silent


def test_per_socket_uses_socket_count_and_linux_lscpu_strings(db):
    ent(db, product="Oracle Database", publisher="Oracle", license_metric="per_socket", total_quantity=4)
    asset(db, "win", [sw("Oracle Database 19c", publisher="Oracle Corporation")],
          platform_properties={"cpu": {"data": {"sockets": 2, "physical_cores": 16}}})
    asset(db, "lnx", [sw("Oracle Database 19c", publisher="Oracle Corporation")],
          platform_properties={"cpu": {"data": {"sockets": "2", "cores_per_socket": "8"}}})
    assert pos(db, "Oracle Database")["consumed"] == 4


def test_per_user_named_counts_distinct_users_not_devices(db):
    ent(db, product="Office", license_metric="per_user_named", total_quantity=2)
    asset(db, "pc-a1", [sw(OFFICE, publisher="Microsoft")], assigned_user="CORP\\alice")
    asset(db, "pc-a2", [sw(OFFICE, publisher="Microsoft")], assigned_user="corp\\ALICE")
    asset(db, "pc-b", [sw(OFFICE, publisher="Microsoft")], assigned_user="CORP\\bob")
    p = pos(db, "Office")
    assert p["consumed"] == 2 and p["position"] == 0 and p["status"] == "compliant"
    assert sum(1 for a in p["consuming_assets"] if a["quantity"] == 0) == 1     # alice's second device


def test_unmeasurable_metric_is_never_reported_as_shelfware_or_over_deployed(db):
    ent(db, publisher="Microsoft", product="Windows Server CAL", license_metric="per_cal", total_quantity=500,
        unit_cost=50.0)
    asset(db, "pc0", [sw("Microsoft Windows Server CAL", publisher="Microsoft")])
    p = pos(db, "Windows Server CAL")
    assert p["status"] == "not_measured" and p["reclaimable_count"] == 0 and p["exposure_cost"] == 0.0
    s = S.compute_summary(db, TID)
    assert s["tracked_products"] == 0 and s["not_measured_products"] == 1 and s["reclaimable_count"] == 0
    assert S.compute_license_positions(db, TID)["unlicensed"] == []       # recorded, not flagged


def test_open_source_is_unlimited_and_never_over_deployed(db):
    ent(db, publisher="Igor Pavlov", product="7-Zip", license_type="open_source",
        license_metric="per_installation", total_quantity=1)
    for i in range(4):
        asset(db, f"pc{i}", [sw("7-Zip 23.01 (x64)", publisher="Igor Pavlov")])
    p = pos(db, "7-Zip")
    assert p["consumed"] == 4 and p["position"] >= 0 and p["status"] == "compliant"


def test_explicit_software_key_beats_name_matching(db):
    ent(db, publisher="Microsoft", product="SQL Server", software_key="mssql-2022", total_quantity=1)
    asset(db, "db0", [sw("Microsoft SQL Server 2022 (64-bit)", publisher="Microsoft", key="mssql-2022")])
    asset(db, "db1", [sw("Microsoft SQL Server 2019 (64-bit)", publisher="Microsoft", key="mssql-2019")])
    assert pos(db, "SQL Server")["consumed"] == 1
    assert [u["software_key"] for u in S.compute_license_positions(db, TID)["unlicensed"]] == ["mssql-2019"]


# ── reclaim proxy ───────────────────────────────────────────────────────────

def test_stale_and_decommissioned_hosts_are_reclaim_candidates(db):
    ent(db, total_quantity=3, unit_cost=100.0)
    asset(db, "live", [sw(OFFICE, publisher="Microsoft")])
    stale = asset(db, "stale", [sw(OFFICE, publisher="Microsoft")])
    stale.last_seen_at = NOW - timedelta(days=120)
    asset(db, "gone", [sw(OFFICE, publisher="Microsoft")], lifecycle_state="decommissioned")
    db.commit()
    p = pos(db, "Office")
    assert p["consumed"] == 3 and p["reclaimable_count"] == 2          # stale + decommissioned
    flagged = {r.asset_id for r in db.query(LicenseAssignment).filter(LicenseAssignment.is_reclaim_candidate.is_(True))}
    assert flagged == {stale.id, db.query(ITAsset).filter_by(name="gone").one().id}
    assert S.compute_summary(db, TID)["reclaimable_savings"] == 200.0


def test_never_seen_host_is_not_called_stale(db):
    ent(db, total_quantity=1)
    a = asset(db, "x", [sw(OFFICE, publisher="Microsoft")])
    a.last_seen_at = None
    db.commit()
    assert pos(db, "Office")["reclaimable_count"] == 0


# ── rebuild behaviour ───────────────────────────────────────────────────────

def test_rebuild_is_idempotent_and_write_on_change(db):
    ent(db, total_quantity=5)
    asset(db, "pc0", [sw(OFFICE, publisher="Microsoft")])
    first = S.rebuild_assignments(db, TID)
    assert first["inserted"] == 1
    again = S.rebuild_assignments(db, TID)
    assert (again["inserted"], again["updated"], again["deleted"]) == (0, 0, 0)


def test_software_removed_from_asset_removes_its_seat(db):
    ent(db, total_quantity=5)
    a = asset(db, "pc0", [sw(OFFICE, publisher="Microsoft")])
    assert pos(db, "Office")["consumed"] == 1
    a.detected_software_json = []
    db.commit()
    assert pos(db, "Office")["consumed"] == 0


def test_deleting_an_entitlement_releases_its_installs_to_unlicensed(db):
    e = ent(db, total_quantity=5)
    asset(db, "pc0", [sw(OFFICE, publisher="Microsoft")])
    assert pos(db, "Office")["consumed"] == 1
    assert S.delete_entitlement(db, TID, e.id) is True
    out = S.compute_license_positions(db, TID)
    assert out["positions"] == [] and out["unlicensed"][0]["install_count"] == 1
    assert S.delete_entitlement(db, TID, e.id) is False


def test_manual_assignments_survive_rebuild_and_block_double_count(db):
    e = ent(db, total_quantity=5)
    a = asset(db, "pc0", [sw(OFFICE, publisher="Microsoft", key="office-key")])
    db.add(LicenseAssignment(tenant_id=TID, entitlement_id=e.id, asset_id=a.id, software_key="manual-key",
                             quantity_consumed=1, assignment_source="manual"))
    db.commit()
    S.rebuild_assignments(db, TID)
    rows = db.query(LicenseAssignment).all()
    assert [r.assignment_source for r in rows] == ["manual"]            # no auto duplicate for that asset+pool
    assert pos(db, "Office")["consumed"] == 1


def test_ad_hoc_pentest_targets_and_application_assets_are_not_inventory(db):
    ent(db, total_quantity=5)
    asset(db, "adhoc", [sw(OFFICE, publisher="Microsoft")], status="adhoc")
    asset(db, "app", [sw(OFFICE, publisher="Microsoft")], asset_role="application")
    asset(db, "real", [sw(OFFICE, publisher="Microsoft")])
    assert pos(db, "Office")["consumed"] == 1


def test_refresh_on_read_is_debounced_by_ttl(db, monkeypatch):
    monkeypatch.setattr(S, "REBUILD_TTL_SECONDS", 3600)
    calls = []
    real = S._rebuild_locked
    monkeypatch.setattr(S, "_rebuild_locked", lambda d, t: calls.append(1) or real(d, t))
    ent(db, total_quantity=5)
    asset(db, "pc0", [sw(OFFICE, publisher="Microsoft")])
    S.compute_summary(db, TID)
    S.compute_summary(db, TID)
    S.compute_license_positions(db, TID)
    assert len(calls) == 1
    S.refresh_if_stale(db, TID, force=True)
    assert len(calls) == 2


def test_a_failing_rebuild_never_breaks_a_read(db, monkeypatch):
    ent(db, total_quantity=1)
    monkeypatch.setattr(S, "_rebuild_locked", lambda d, t: (_ for _ in ()).throw(RuntimeError("boom")))
    assert S.compute_summary(db, TID)["entitlement_lines"] == 1


# ── summary ─────────────────────────────────────────────────────────────────

def test_empty_tenant_is_honest_zeros_and_null_percent(db):
    s = S.compute_summary(db, TID)
    assert s["compliance_pct"] is None and s["tracked_products"] == 0
    for k in ("compliant", "over_deployed", "under_licensed", "unlicensed_count", "expiring_30", "expiring_60",
              "expiring_90", "annual_spend", "renewal_cost_90d", "reclaimable_count", "reclaimable_savings"):
        assert s[k] == 0, k
    assert S.compute_license_positions(db, TID) == {"positions": [], "unlicensed": []}
    assert S.list_entitlements(db, TID) == {"entitlements": []}


def test_summary_numbers_and_invariant(db):
    # product A: compliant, B: over-deployed (2 short @ 10), C: expiring in 20 days, D unlicensed install
    ent(db, publisher="Acme", product="Alpha", total_quantity=2)
    ent(db, publisher="Acme", product="Beta", total_quantity=1, unit_cost=10.0)
    ent(db, publisher="Acme", product="Gamma", total_quantity=1, expiry_date=NOW + timedelta(days=20),
        license_type="subscription", unit_cost=100.0, start_date=NOW - timedelta(days=345))
    for i in range(2):
        asset(db, f"a{i}", [sw("Acme Alpha", publisher="Acme Inc")])
    for i in range(3):
        asset(db, f"b{i}", [sw("Acme Beta", publisher="Acme Inc")])
    asset(db, "g", [sw("Acme Gamma", publisher="Acme Inc"), sw("Mystery Tool", publisher="Nobody")])
    s = S.compute_summary(db, TID)
    assert s["tracked_products"] == 3 and s["over_deployed"] == 1 and s["compliant"] == 2
    assert s["compliant"] + s["over_deployed"] == s["tracked_products"]
    assert s["under_licensed"] == s["over_deployed"]
    assert s["compliance_pct"] == 66.7
    assert s["unlicensed_count"] == 1
    assert (s["expiring_30"], s["expiring_60"], s["expiring_90"]) == (1, 1, 1)
    assert s["exposure_cost"] == 20.0 and s["licenses_short"] == 2
    assert s["top_offenders"][0]["product"] == "Beta"
    # Gamma: 1 seat x 100, 365-day term (start..expiry end) -> annual ~100, renewal 100
    assert s["annual_spend"] == pytest.approx(100.0, abs=1.0) and s["renewal_cost_90d"] == 100.0
    assert s["currency"] == "USD" and s["assets_with_inventory"] == 6


def test_cumulative_expiry_windows_include_maintenance_expiry(db):
    ent(db, product="A", total_quantity=1, expiry_date=NOW + timedelta(days=25))
    ent(db, product="B", total_quantity=1, expiry_date=NOW + timedelta(days=55))
    ent(db, product="C", total_quantity=1, expiry_date=NOW + timedelta(days=85))
    ent(db, product="D", total_quantity=1, expiry_date=NOW + timedelta(days=400))
    ent(db, product="E", total_quantity=1, maintenance_expiry=NOW + timedelta(days=10))      # perpetual + SA
    s = S.compute_summary(db, TID)
    assert (s["expiring_30"], s["expiring_60"], s["expiring_90"]) == (2, 3, 4)
    assert [n["product"] for n in s["nearest_expiries"]][:2] == ["E", "A"]


def test_perpetual_is_not_annual_spend_and_mixed_currencies_are_not_summed(db):
    ent(db, product="Perp", total_quantity=10, unit_cost=100.0)                         # perpetual: one-time
    ent(db, product="SubUSD", license_type="subscription", total_quantity=10, unit_cost=10.0, currency="USD",
        expiry_date=NOW + timedelta(days=200))
    ent(db, product="SubEUR", license_type="subscription", total_quantity=1, unit_cost=5.0, currency="EUR",
        expiry_date=NOW + timedelta(days=200))
    s = S.compute_summary(db, TID)
    assert s["currency"] == "USD" and s["annual_spend"] == 100.0
    assert s["by_currency"]["EUR"]["annual_spend"] == 5.0


# ── register + asset view ───────────────────────────────────────────────────

def test_register_lines_carry_live_numbers_and_status_vocabulary(db):
    ent(db, total_quantity=1, product_key=None)
    ent(db, product="Old", total_quantity=3, expiry_date=NOW - timedelta(days=5))
    gone = ent(db, product="Gone", total_quantity=3)
    gone.status = "retired"
    db.commit()
    asset(db, "pc0", [sw(OFFICE, publisher="Microsoft")])
    asset(db, "pc1", [sw(OFFICE, publisher="Microsoft")])
    by = {l["product"]: l for l in S.list_entitlements(db, TID)["entitlements"]}
    assert (by["Office"]["consumed"], by["Office"]["position"], by["Office"]["status"]) == (2, -1, "over_deployed")
    assert (by["Old"]["status"], by["Old"]["consumed"], by["Old"]["position"]) == ("expired", None, None)
    assert (by["Gone"]["status"], by["Gone"]["consumed"]) == ("retired", None)
    for k in ("id", "publisher", "product", "edition", "license_type", "license_metric", "total_quantity",
              "consumed", "position", "status", "unit_cost", "currency", "purchase_date", "start_date",
              "expiry_date", "maintenance_expiry", "po_reference", "contract_reference", "reseller", "notes"):
        assert k in by["Office"], k


def test_asset_licenses_statuses(db):
    ent(db, product="Office", total_quantity=1)
    ent(db, publisher="Acme", product="Soon", total_quantity=5, expiry_date=NOW + timedelta(days=20))
    a = asset(db, "pc0", [sw(OFFICE, publisher="Microsoft"), sw("Acme Soon", publisher="Acme"),
                          sw("Mystery Tool", publisher="Nobody")])
    asset(db, "pc1", [sw(OFFICE, publisher="Microsoft")])
    out = S.asset_licenses(db, TID, a.id)
    by = {l["detected_name"]: l for l in out["licenses"]}
    assert by[OFFICE]["status"] == "over_deployed" and by[OFFICE]["entitlement_product"] == "Office"
    assert by["Acme Soon"]["status"] == "expiring"
    assert by["Mystery Tool"]["status"] == "unlicensed" and by["Mystery Tool"]["entitlement_id"] is None
    assert by["Mystery Tool"]["publisher"] == "Nobody" and by["Acme Soon"]["publisher"] == "Acme"
    assert [l["status"] for l in out["licenses"]][0] == "unlicensed"        # worst first
    assert out["summary"] == {"licensed": 0, "unlicensed": 1, "over_deployed": 1, "expiring": 1, "total": 3}
    assert S.asset_licenses(db, TID, 9999) is None
    assert S.asset_licenses(db, 2, a.id) is None                           # other tenant


# ── CRUD + validation ───────────────────────────────────────────────────────

def test_clean_entitlement_validation_and_coercion():
    f = S.clean_entitlement({"publisher": " MS ", "product": "X", "total_quantity": "1,200",
                             "unit_cost": "$1,234.50", "license_type": "Volume License", "license_metric": "Per Core",
                             "expiry_date": "2026-12-31", "currency": "usd"})
    assert f["total_quantity"] == 1200 and f["unit_cost"] == 1234.5 and f["publisher"] == "MS"
    assert (f["license_type"], f["license_metric"], f["currency"]) == ("volume", "per_core", "USD")
    assert f["expiry_date"] == datetime(2026, 12, 31)
    for bad in ({"product": "x", "total_quantity": 1},                                   # no publisher
                {"publisher": "p", "product": "x", "total_quantity": -1},
                {"publisher": "p", "product": "x", "total_quantity": 1.5},
                {"publisher": "p", "product": "x"},                                       # no quantity
                {"publisher": "p", "product": "x", "total_quantity": 1, "license_type": "weird"},
                {"publisher": "p", "product": "x", "total_quantity": 1, "license_metric": "weird"},
                {"publisher": "p", "product": "x", "total_quantity": 1, "unit_cost": "abc"},
                {"publisher": "p", "product": "x", "total_quantity": 1, "expiry_date": "soon"}):
        with pytest.raises(ValueError):
            S.clean_entitlement(bad)


def test_register_echoed_computed_fields_are_ignored():
    f = S.clean_entitlement({"publisher": "p", "product": "x", "total_quantity": 1, "consumed": 9,
                             "position": -3, "status": "over_deployed", "id": 77, "created_at": "x"})
    assert "consumed" not in f and "position" not in f and "id" not in f and "status" not in f


def test_expiry_before_start_rejected(db):
    with pytest.raises(ValueError):
        S.create_entitlement(db, TID, {"publisher": "p", "product": "x", "total_quantity": 1,
                                       "start_date": "2026-06-01", "expiry_date": "2026-01-01"})


def test_put_only_changes_the_keys_sent(db):
    e = S.create_entitlement(db, TID, {"publisher": "p", "product": "x", "total_quantity": 1, "software_key": "k1",
                                       "notes": "n"}, rebuild=False)
    S.update_entitlement(db, TID, e.id, {"total_quantity": 7, "notes": None})
    db.refresh(e)
    assert (e.total_quantity, e.software_key, e.notes) == (7, "k1", None)
    assert S.update_entitlement(db, TID, 9999, {"total_quantity": 1}) is None
    assert S.update_entitlement(db, 2, e.id, {"total_quantity": 1}) is None          # other tenant


def test_extending_expiry_revives_an_expired_line_but_resaving_does_not(db):
    e = S.create_entitlement(db, TID, {"publisher": "p", "product": "x", "total_quantity": 1,
                                       "expiry_date": "2026-01-01"}, rebuild=False)
    S.sweep_expired(db, TID)
    db.refresh(e)
    assert e.status == "expired"
    S.update_entitlement(db, TID, e.id, {"notes": "just a note", "expiry_date": "2026-01-01"})
    db.refresh(e)
    assert e.status == "expired"
    S.update_entitlement(db, TID, e.id, {"expiry_date": "2027-06-30"})
    db.refresh(e)
    assert e.status == "active"


def test_product_key_is_encrypted_at_rest_and_never_returned(db, monkeypatch):
    monkeypatch.setenv("SESSION_SECRET", "unit-test-secret")
    e = S.create_entitlement(db, TID, {"publisher": "p", "product": "x", "total_quantity": 1,
                                       "product_key": "ABCDE-FGHIJ-KLMNO"}, rebuild=False)
    assert e.product_key.startswith("enc:v1:") and "ABCDE" not in e.product_key
    line = S.list_entitlements(db, TID)["entitlements"][0]
    assert line["has_product_key"] is True and "product_key" not in line
    assert "ABCDE" not in repr(line)
    S.update_entitlement(db, TID, e.id, {"product_key": ""})
    db.refresh(e)
    assert e.product_key is None


def test_product_key_is_refused_not_stored_plain_when_encryption_is_unavailable(db, monkeypatch):
    monkeypatch.delenv("SESSION_SECRET", raising=False)
    with pytest.raises(S.EncryptionUnavailable):
        S.create_entitlement(db, TID, {"publisher": "p", "product": "x", "total_quantity": 1,
                                       "product_key": "SECRET"}, rebuild=False)
    assert db.query(LicenseEntitlement).count() == 0


# ── import ──────────────────────────────────────────────────────────────────

CSV = (
    "Vendor,Software,Qty,Price,PO Number,Expiration Date,License Type,Metric\n"
    "Microsoft,Microsoft 365 Apps,250,180.00,PO-1,2027-01-31,Subscription,Per User\n"
    "Adobe,Acrobat Pro,10,239.99,PO-2,,perpetual,per device\n"
    "Adobe,Bad Row,abc,1,PO-3,,perpetual,per device\n"
)


def test_csv_import_maps_synonym_headers_validates_rows_and_is_idempotent(db):
    dry = importer.import_entitlements(db, TID, CSV.encode(), "po.csv", dry_run=True)
    assert dry["would_create"] == 2 and dry["error_count"] == 1 and db.query(LicenseEntitlement).count() == 0
    assert dry["errors"][0]["row"] == 4 and dry["columns"]["Vendor"] == "publisher"

    out = importer.import_entitlements(db, TID, CSV.encode(), "po.csv", user="u@x")
    assert out["created"] == 2 and out["skipped"] == 0 and len(out["ids"]) == 2
    rows = {r.product: r for r in db.query(LicenseEntitlement)}
    assert rows["Microsoft 365 Apps"].license_metric == "per_user_named"
    assert rows["Microsoft 365 Apps"].license_type == "subscription"
    assert rows["Acrobat Pro"].unit_cost == 239.99 and rows["Acrobat Pro"].created_by == "u@x"

    again = importer.import_entitlements(db, TID, CSV.encode(), "po.csv")
    assert again["created"] == 0 and again["skipped"] == 2 and db.query(LicenseEntitlement).count() == 2


def test_import_rejects_unusable_files(db):
    with pytest.raises(ValueError):
        importer.import_entitlements(db, TID, b"", "x.csv")
    with pytest.raises(ValueError):
        importer.import_entitlements(db, TID, b"colA,colB\n1,2\n", "x.csv")           # no publisher/product columns
    with pytest.raises(ValueError):
        importer.import_entitlements(db, TID, b"a,b", "legacy.xls")


def test_xlsx_import(db):
    from io import BytesIO
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Publisher", "Product", "Quantity", "Expiry"])
    ws.append(["Zoom", "Zoom Workplace", 25, datetime(2027, 3, 1)])
    buf = BytesIO()
    wb.save(buf)
    out = importer.import_entitlements(db, TID, buf.getvalue(), "licenses.xlsx")
    assert out["created"] == 1
    assert db.query(LicenseEntitlement).one().expiry_date == datetime(2027, 3, 1)


def test_template_round_trips_through_the_importer(db):
    out = importer.import_entitlements(db, TID, importer.template_csv().encode(), "t.csv")
    assert out["created"] == 2 and out["error_count"] == 0


# ── wiring ──────────────────────────────────────────────────────────────────

def test_router_literal_routes_precede_the_id_catch_alls():
    from grc.modules.license_management.router import router
    paths = [(sorted(r.methods)[0], r.path) for r in router.routes]
    idx = {p: i for i, p in enumerate(paths)}
    catch_all = min(idx[("DELETE", "/licenses/{license_id}")], idx[("PUT", "/licenses/{license_id}")])
    for literal in (("GET", "/licenses/positions"), ("GET", "/licenses/summary"),
                    ("GET", "/licenses/asset/{asset_id}"), ("POST", "/licenses/import"),
                    ("GET", "/licenses/import/template")):
        assert idx[literal] < catch_all, literal


def test_router_contract_endpoints_exist():
    from grc.modules.license_management.router import router
    have = {(m, r.path) for r in router.routes for m in r.methods}
    for need in (("GET", "/licenses"), ("POST", "/licenses"), ("PUT", "/licenses/{license_id}"),
                 ("DELETE", "/licenses/{license_id}"), ("GET", "/licenses/positions"),
                 ("GET", "/licenses/summary"), ("GET", "/licenses/asset/{asset_id}"),
                 ("POST", "/licenses/import")):
        assert need in have, need


def test_router_is_mounted_in_main():
    import inspect
    import grc.main as main
    src = inspect.getsource(main)
    assert "grc.modules.license_management.router import router as license_management_router" in src
    assert "app.include_router(license_management_router)" in src
    assert any(getattr(r, "path", "") == "/licenses/positions" for r in main.app.routes)


def test_models_registered_on_metadata_with_expected_tables():
    assert "grc_license_entitlements" in Base.metadata.tables
    assert "grc_license_assignments" in Base.metadata.tables
    cols = {c.name for c in LicenseEntitlement.__table__.columns}
    assert {"publisher", "product", "edition", "license_type", "license_metric", "total_quantity", "unit_cost",
            "currency", "po_reference", "contract_reference", "reseller", "purchase_date", "start_date",
            "expiry_date", "maintenance_included", "maintenance_expiry", "support_level", "downgrade_rights",
            "product_key", "software_key", "catalog_product_id", "status", "notes", "created_at", "updated_at",
            "created_by", "tenant_id", "id"} <= cols
    acols = {c.name for c in LicenseAssignment.__table__.columns}
    assert {"id", "tenant_id", "entitlement_id", "asset_id", "user_identifier", "software_key", "detected_name",
            "detected_version", "quantity_consumed", "assignment_source", "is_reclaim_candidate", "assigned_at",
            "assigned_by", "notes"} <= acols


def test_assignment_foreign_keys_cascade_so_asset_delete_code_needs_no_change():
    fks = {fk.parent.name: fk.ondelete for fk in LicenseAssignment.__table__.foreign_keys}
    assert fks == {"entitlement_id": "CASCADE", "asset_id": "CASCADE"}


def test_report_dataset_registered_and_every_field_is_a_real_column():
    from grc.routers import reporting_router as R
    spec = R.SERVER_DATASETS["licenses"]
    assert "licenses" in R.AVA_DATASET_ALLOWLIST and R.DATASET_LABELS["licenses"]
    cols = {c.name for c in spec.model.__table__.columns}
    for f in set(spec.fields) | set(spec.sortable) | set(spec.filterable) | set(spec.search):
        assert f in cols, f
    assert "product_key" not in spec.fields
    assert spec.permissions == ["assets:asset_inventory:*"]
    assert spec.default_order[0] in cols


# ── pool unification + match precedence ─────────────────────────────────────

def test_publisher_prefix_in_product_name_does_not_split_a_pool(db):
    ent(db, product="Microsoft Office Professional Plus 2019", total_quantity=1)
    ent(db, product="Office Professional Plus 2019", total_quantity=1, po_reference="PO-2")
    for i in range(3):
        asset(db, f"pc{i}", [sw(OFFICE, publisher="Microsoft")])
    out = S.compute_license_positions(db, TID)["positions"]
    assert len(out) == 1 and (out[0]["entitled"], out[0]["consumed"]) == (2, 3)


def test_a_stronger_match_on_an_expired_pool_does_not_hide_a_live_pool(db):
    ent(db, publisher="Microsoft", product="SQL Server", software_key="mssql-2022", total_quantity=5,
        expiry_date=NOW - timedelta(days=30))                       # tier-1 key match, but lapsed
    ent(db, publisher="Microsoft", product="SQL Server 2022", total_quantity=1)       # tier-3 name match, live
    asset(db, "db0", [sw("Microsoft SQL Server 2022 (64-bit)", publisher="Microsoft", key="mssql-2022")])
    live = [p for p in S.compute_license_positions(db, TID)["positions"] if p["product"] == "SQL Server 2022"][0]
    assert live["consumed"] == 1 and live["status"] == "compliant"


def test_explicit_key_ignores_publisher_but_derived_slug_key_does_not():
    d = M.Descriptor("Adobe", "Reader", None, [])
    foxit = M.read_install({"name": "Reader", "publisher": "Foxit Software", "source": "registry"})
    assert d.tier(foxit) is None
    d2 = M.Descriptor("Adobe", "Anything", None, ["reader"])
    assert d2.tier(foxit) == 1


def test_parse_date_accepts_iso_and_unambiguous_text_and_rejects_the_rest():
    assert S.parse_date("2026-12-31") == datetime(2026, 12, 31)
    assert S.parse_date("2026-12-31T10:30:00Z") == datetime(2026, 12, 31, 10, 30)
    assert S.parse_date("2026-12-31T10:30:00+02:00") == datetime(2026, 12, 31, 8, 30)       # naive UTC
    assert S.parse_date("31-Dec-2026") == datetime(2026, 12, 31)
    assert S.parse_date("") is None and S.parse_date(None) is None and S.parse_date("Perpetual") is None
    with pytest.raises(ValueError):
        S.parse_date("12/31/2026")


def test_enum_inputs_are_forgiving_about_case_spacing_and_parentheses():
    f = S.clean_entitlement({"publisher": "p", "product": "x", "total_quantity": 1,
                             "license_type": "Open Source", "license_metric": "Per user (named)"})
    assert (f["license_type"], f["license_metric"]) == ("open_source", "per_user_named")
    f = S.clean_entitlement({"publisher": "p", "product": "x", "total_quantity": 1, "license_metric": "per-core"})
    assert f["license_metric"] == "per_core"


# ── HTTP contract (FastAPI TestClient over the SQLite session) ──────────────

@pytest.fixture()
def client(db):
    from types import SimpleNamespace
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from grc.models import Tenant, get_db
    from grc.routers.auth_router import require_auth
    from grc.modules.license_management import router as lr

    Tenant.__table__.create(db.get_bind(), checkfirst=True)
    db.add(Tenant(id=TID, name="T", slug="t"))
    db.commit()
    app = FastAPI()
    app.include_router(lr.router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[require_auth] = lambda: SimpleNamespace(id=1, email="admin@ava.io", username="admin")
    for gate in (lr._can_create, lr._can_edit, lr._can_delete):     # the signed-in user may write
        app.dependency_overrides[gate] = lambda: True
    app.state.gates = (lr._can_create, lr._can_edit, lr._can_delete)
    return TestClient(app)


BODY = {"publisher": "Microsoft", "product": "Office Professional Plus 2019", "edition": "", "license_type": "perpetual",
        "license_metric": "per_device", "total_quantity": 2, "unit_cost": 150, "currency": "USD",
        "purchase_date": "2026-01-01", "start_date": None, "expiry_date": "", "maintenance_expiry": None,
        "po_reference": "PO-1", "contract_reference": None, "reseller": None, "notes": None}


def test_http_crud_roundtrip_and_error_shapes(client, db):
    r = client.post("/licenses", json=BODY)
    assert r.status_code == 200 and set(r.json()) == {"id"}
    lid = r.json()["id"]

    r = client.get("/licenses")
    assert r.status_code == 200
    line = r.json()["entitlements"][0]
    for k in ("id", "publisher", "product", "edition", "license_type", "license_metric", "total_quantity",
              "consumed", "position", "status", "unit_cost", "currency", "purchase_date", "start_date",
              "expiry_date", "maintenance_expiry", "po_reference", "contract_reference", "reseller", "notes"):
        assert k in line, k
    assert line["edition"] is None and line["purchase_date"] == "2026-01-01T00:00:00"

    # the register echoes computed fields back on save: they must be ignored, not rejected
    echoed = {**BODY, "total_quantity": 9, "consumed": 5, "position": -4, "status": "over_deployed", "id": lid}
    assert client.put(f"/licenses/{lid}", json=echoed).json() == {"id": lid}
    assert client.get("/licenses").json()["entitlements"][0]["total_quantity"] == 9

    bad = client.post("/licenses", json={**BODY, "publisher": " "})
    assert bad.status_code == 400 and bad.json()["detail"] == "publisher is required."
    assert client.post("/licenses", json={**BODY, "license_metric": "per_banana"}).status_code == 400
    assert client.put("/licenses/99999", json={"notes": "x"}).status_code == 404

    assert client.delete(f"/licenses/{lid}").json() == {"id": lid, "deleted": True}
    assert client.delete(f"/licenses/{lid}").status_code == 404
    assert client.get("/licenses").json() == {"entitlements": []}


def test_http_positions_summary_and_asset_shapes(client, db):
    client.post("/licenses", json={**BODY, "total_quantity": 1})
    a = asset(db, "pc0", [sw(OFFICE, publisher="Microsoft"), sw("Mystery Tool", publisher="Nobody")])
    asset(db, "pc1", [sw(OFFICE, publisher="Microsoft")])

    pos = client.get("/licenses/positions").json()
    assert set(pos) == {"positions", "unlicensed"}
    p = pos["positions"][0]
    for k in ("publisher", "product", "edition", "license_metric", "entitled", "consumed", "position", "status",
              "unit_cost", "exposure_cost", "reclaimable_count", "nearest_expiry", "consuming_assets"):
        assert k in p, k
    assert set(p["consuming_assets"][0]) == {"asset_id", "name", "quantity"}
    u = pos["unlicensed"][0]
    assert {"software_key", "name", "publisher", "install_count", "assets"} <= set(u)
    assert set(u["assets"][0]) == {"asset_id", "name"}

    s = client.get("/licenses/summary").json()
    for k in ("compliance_pct", "tracked_products", "compliant", "over_deployed", "under_licensed",
              "unlicensed_count", "expiring_30", "expiring_60", "expiring_90", "annual_spend",
              "renewal_cost_90d", "reclaimable_count", "reclaimable_savings"):
        assert k in s, k
    assert (s["tracked_products"], s["over_deployed"], s["unlicensed_count"]) == (1, 1, 1)

    asset_view = client.get(f"/licenses/asset/{a.id}").json()
    assert set(asset_view["licenses"][0]) >= {"detected_name", "detected_version", "publisher", "software_key",
                                              "status", "entitlement_id", "entitlement_product"}
    assert {l["status"] for l in asset_view["licenses"]} == {"over_deployed", "unlicensed"}
    assert client.get("/licenses/asset/424242").status_code == 404


def test_http_import_template_and_upload(client, db):
    t = client.get("/licenses/import/template")
    assert t.status_code == 200 and t.text.splitlines()[0].startswith("publisher,product")
    files = {"file": ("po.csv", CSV.encode(), "text/csv")}
    dry = client.post("/licenses/import?dry_run=true", files=files).json()
    assert dry["dry_run"] is True and dry["would_create"] == 2 and dry["created"] == 0
    done = client.post("/licenses/import", files={"file": ("po.csv", CSV.encode(), "text/csv")}).json()
    assert done["created"] == 2 and done["error_count"] == 1 and len(done["ids"]) == 2
    assert {"columns", "unmapped", "errors", "skipped", "total"} <= set(done)
    bad = client.post("/licenses/import", files={"file": ("x.csv", b"", "text/csv")})
    assert bad.status_code == 400


def test_writes_are_permission_gated_but_reads_only_need_a_signed_in_user(client, db):
    from fastapi.routing import APIRoute
    from grc.modules.license_management import router as lr
    gated = {}
    for r in lr.router.routes:
        assert isinstance(r, APIRoute)
        deps = {d.call for d in r.dependant.dependencies}
        for m in r.methods:
            gated[(m, r.path)] = deps
    assert lr._can_create in gated[("POST", "/licenses")]
    assert lr._can_create in gated[("POST", "/licenses/import")]
    assert lr._can_edit in gated[("PUT", "/licenses/{license_id}")]
    assert lr._can_delete in gated[("DELETE", "/licenses/{license_id}")]
    for read in (("GET", "/licenses"), ("GET", "/licenses/positions"), ("GET", "/licenses/summary"),
                 ("GET", "/licenses/asset/{asset_id}")):
        assert not ({lr._can_create, lr._can_edit, lr._can_delete} & gated[read]), read


def test_unauthenticated_writes_are_rejected_before_anything_is_stored(client, db):
    for gate in client.app.state.gates:
        client.app.dependency_overrides.pop(gate, None)           # real permission check, no token
    assert client.post("/licenses", json=BODY).status_code == 401
    assert client.put("/licenses/1", json=BODY).status_code == 401
    assert client.delete("/licenses/1").status_code == 401
    assert db.query(LicenseEntitlement).count() == 0


def test_product_key_never_reaches_the_audit_trail():
    from grc import audit_logger, rich_audit
    body = {"publisher": "p", "product_key": "ABCDE-FGHIJ", "nested": [{"License_Key": "X"}]}
    clean = audit_logger._sanitize_value(body)
    assert clean["product_key"] == "***" and clean["nested"][0]["License_Key"] == "***"
    assert rich_audit._sanitize(body)["product_key"] == "***"


def test_xlsx_import_refuses_non_workbooks_and_oversized_sheets(db, monkeypatch):
    with pytest.raises(ValueError, match="valid .xlsx"):
        importer.import_entitlements(db, TID, b"this is not a zip", "licenses.xlsx")
    from io import BytesIO
    from openpyxl import Workbook
    monkeypatch.setattr(importer, "MAX_ROWS", 3)
    wb = Workbook()
    ws = wb.active
    ws.append(["Publisher", "Product", "Quantity"])
    for i in range(300):
        ws.append(["P", f"Product {i}", 1])
    buf = BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="Too many rows"):
        importer.import_entitlements(db, TID, buf.getvalue(), "licenses.xlsx")
    assert db.query(LicenseEntitlement).count() == 0


def test_a_product_with_no_install_on_record_is_never_called_shelfware(db):
    """OS licences, SaaS tenants and unscanned estates have no install footprint: the
    inventory proves nothing either way, so no surplus / savings may be claimed."""
    ent(db, publisher="Microsoft", product="Windows Server 2022", edition="Standard", license_metric="per_core",
        total_quantity=16, unit_cost=100.0)
    ent(db, publisher="Salesforce", product="Sales Cloud", license_type="saas", license_metric="per_user_named",
        total_quantity=100, unit_cost=1500.0)
    asset(db, "pc0", [sw("Some Unrelated App", publisher="Nobody")])
    rows = {p["product"]: p for p in S.compute_license_positions(db, TID)["positions"]}
    for name in ("Windows Server 2022", "Sales Cloud"):
        assert rows[name]["status"] == "not_detected" and rows[name]["reclaimable_count"] == 0, name
        assert rows[name]["consuming_assets"] == []
    s = S.compute_summary(db, TID)
    assert s["reclaimable_count"] == 0 and s["reclaimable_savings"] == 0.0
    assert s["tracked_products"] == 0 and s["compliance_pct"] is None        # no fake 100%
    assert s["not_detected_products"] == 2 and s["compliant"] + s["over_deployed"] == s["tracked_products"]
    lines = S.list_entitlements(db, TID)["entitlements"]
    assert {l["status"] for l in lines} == {"not_detected"}


def test_an_undetected_product_that_is_about_to_expire_still_says_so(db):
    ent(db, product="Sales Cloud", publisher="Salesforce", license_type="saas", total_quantity=5,
        expiry_date=NOW + timedelta(days=20))
    assert pos(db, "Sales Cloud")["status"] == "expiring_soon"


def test_detection_flips_the_same_product_to_tracked(db):
    ent(db, total_quantity=1)
    assert pos(db, "Office")["status"] == "not_detected"
    asset(db, "pc0", [sw(OFFICE, publisher="Microsoft")])
    s = S.compute_summary(db, TID)
    assert pos(db, "Office")["status"] == "compliant" and s["tracked_products"] == 1 and s["compliance_pct"] == 100.0
