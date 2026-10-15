"""Scheduled & emailed reports (grc/modules/report_scheduler).

Runs against an in-memory SQLite DB holding only the tables the feature touches, the
REAL Report Builder export path (reporting_router.export_dataset) and a stubbed
smtplib - no network, no Postgres. Covers: the schedule maths, input validation (incl.
the "an unusable filter must fail the run, not broaden it" rule), file generation in
every format, SMTP configuration order, the email itself, failure recording, the
exactly-once claim of due schedules, the HTTP layer's access control and the Celery
fan-out.
"""
import email
import io
from datetime import datetime, timedelta
from types import SimpleNamespace

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import grc.models as m
from grc.modules.license_management.models import LicenseEntitlement
from grc.modules.report_scheduler import mailer, service as S
from grc.modules.report_scheduler.models import ScheduledReport

NOW = datetime(2026, 10, 15, 12, 0, 0)
SMTP_ENV = ("SMTP_HOST", "SMTP_USER", "SMTP_PASSWORD", "SMTP_FROM_EMAIL", "SMTP_PORT")


# ── fixtures ────────────────────────────────────────────────────────────────

@pytest.fixture()
def env(monkeypatch):
    """Session + seeded users/licences. `admin` is the tenant primary contact (all
    permissions); `viewer` has no roles (no dataset permission)."""
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    for cls in (m.Tenant, m.GRCUser, m.Role, m.UserRole, m.Permission, m.RolePermission,
                m.ReportDefinition, ScheduledReport, m.WorkflowEmailConfiguration, LicenseEntitlement):
        cls.__table__.create(bind=engine, checkfirst=True)
    for k in SMTP_ENV:
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setattr(S, "_utcnow", lambda: NOW)
    S._ENSURED.clear()

    Session = sessionmaker(bind=engine, autoflush=False, future=True)
    db = Session()
    tenant = m.Tenant(name="Acme Corp", slug="acme", primary_contact_email="admin@acme.test")
    db.add(tenant)
    db.flush()
    users = {}
    for key, display in (("admin", "Alice Admin"), ("viewer", "Vic Viewer"), ("other", "Olga Other")):
        u = m.GRCUser(username=key, email=f"{key}@acme.test", display_name=display, is_active=True)
        db.add(u)
        users[key] = u
    db.flush()
    for i, (pub, prod, st) in enumerate([("Microsoft", "SQL Server 2022", "active"), ("Adobe", "Acrobat Pro", "active"),
                                         ("Oracle", "Database EE", "expired"), ("Atlassian", "Jira <b>Cloud</b>", "active")]):
        db.add(LicenseEntitlement(tenant_id=tenant.id, publisher=pub, product=prod, license_type="perpetual",
                                  license_metric="per_device", total_quantity=10 + i, currency="USD", status=st))
    db.commit()
    yield SimpleNamespace(db=db, Session=Session, tid=tenant.id, **users)
    db.close()


class FakeSMTP:
    sent: list = []
    logins: list = []
    fail_auth = False
    refuse: dict = {}

    def __init__(self, host, port, timeout=None):
        pass

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def ehlo(self):
        pass

    def starttls(self, context=None):
        pass

    def login(self, user, password):
        if FakeSMTP.fail_auth:
            import smtplib
            raise smtplib.SMTPAuthenticationError(535, b"bad credentials")
        FakeSMTP.logins.append((user, password))

    def sendmail(self, frm, to, raw):
        FakeSMTP.sent.append((frm, list(to), raw))
        return dict(FakeSMTP.refuse)


@pytest.fixture()
def smtp(monkeypatch):
    """Outbound email configured via env + a stub transport."""
    FakeSMTP.sent, FakeSMTP.logins, FakeSMTP.fail_auth, FakeSMTP.refuse = [], [], False, {}
    monkeypatch.setattr(mailer.smtplib, "SMTP", FakeSMTP)
    monkeypatch.setenv("SMTP_HOST", "smtp.env.test")
    monkeypatch.setenv("SMTP_USER", "envuser")
    monkeypatch.setenv("SMTP_PASSWORD", "envpw")
    monkeypatch.setenv("SMTP_FROM_EMAIL", "reports@env.test")
    monkeypatch.setenv("SMTP_PORT", "587")
    return FakeSMTP


def grant(env, user, permission):
    """Give `user` a role holding one permission (the way /auth/me resolves them)."""
    role = m.Role(name=f"role-{permission}", tenant_id=env.tid)
    perm = m.Permission(name=permission, resource=permission.split(":")[0], action=permission.split(":")[-1])
    env.db.add_all([role, perm])
    env.db.flush()
    env.db.add_all([m.RolePermission(role_id=role.id, permission_id=perm.id),
                    m.UserRole(user_id=user.id, role_id=role.id, tenant_id=env.tid)])
    env.db.commit()


def make(env, user=None, **over):
    body = {"name": "Active licences", "dataset": "licenses", "columns": ["publisher", "product", "status"],
            "filters": [{"col": "status", "op": "eq", "value": "active"}], "recipients": ["boss@acme.test"],
            "frequency": "daily", "format": "csv", "start_at": NOW + timedelta(hours=2)}
    body.update(over)
    return S.create_schedule(env.db, user or env.admin, env.tid, body)


# ── schedule maths ──────────────────────────────────────────────────────────

D = datetime


@pytest.mark.parametrize("anchor,freq,after,expected", [
    (D(2026, 1, 1, 8), "daily", D(2026, 1, 1, 9), D(2026, 1, 2, 8)),
    (D(2026, 1, 5, 8), "daily", D(2026, 1, 1), D(2026, 1, 5, 8)),                     # before the anchor
    (D(2026, 1, 1, 8), "daily", D(2026, 1, 3, 8), D(2026, 1, 4, 8)),                   # exactly on a slot
    (D(2026, 1, 1, 8), "daily", D(2026, 1, 11, 9), D(2026, 1, 12, 8)),                 # stalled: ONE catch-up slot
    (D(2026, 1, 5, 8), "weekly", D(2026, 1, 6), D(2026, 1, 12, 8)),
    (D(2026, 1, 15, 8), "monthly", D(2026, 1, 20), D(2026, 2, 15, 8)),
    (D(2026, 1, 31, 10), "monthly", D(2026, 2, 15), D(2026, 2, 28, 10)),               # clamps to month end
    (D(2026, 1, 31, 10), "monthly", D(2026, 2, 28, 10), D(2026, 3, 31, 10)),           # ... without drifting
    (D(2026, 11, 30, 10), "monthly", D(2026, 12, 31), D(2027, 1, 30, 10)),
    (D(2028, 1, 31, 10), "monthly", D(2028, 2, 1), D(2028, 2, 29, 10)),                # leap year
])
def test_next_occurrence(anchor, freq, after, expected):
    assert S.next_occurrence(anchor, freq, after) == expected


def test_default_start_is_next_0600_utc():
    assert S.default_start(D(2026, 5, 1, 5, 0)) == D(2026, 5, 1, 6, 0)
    assert S.default_start(D(2026, 5, 1, 7, 0)) == D(2026, 5, 2, 6, 0)


# ── recipients ──────────────────────────────────────────────────────────────

def test_recipients_are_cleaned_and_deduped():
    assert S.clean_recipients(["a@x.io", "A@x.io", "b@y.com"]) == ["a@x.io", "b@y.com"]
    assert S.clean_recipients("a@x.io, b@y.com; c@z.org\nd@w.net") == ["a@x.io", "b@y.com", "c@z.org", "d@w.net"]


@pytest.mark.parametrize("bad", [["not-an-email"], ["a@x.io>\r\nBcc:evil@x.io"], [], [f"u{i}@x.io" for i in range(26)]])
def test_recipients_rejected(bad):
    with pytest.raises(S.ScheduleError):
        S.clean_recipients(bad)


# ── create / validate ───────────────────────────────────────────────────────

def test_create_inline_schedule(env):
    row = make(env, name="Active <weekly>", frequency="weekly")
    s = S.serialize(env.db, [row], env.admin.id)[0]
    assert s["dataset_label"] == "Software Licenses" and s["is_mine"] is True
    assert row.created_by == env.admin.id and row.created_by_name == "Alice Admin"
    assert row.next_run == NOW + timedelta(hours=2) and s["next_run"].endswith("Z")


@pytest.mark.parametrize("over,why", [
    ({"filters": [{"col": "nope", "op": "eq", "value": "1"}]}, "unknown filter column"),
    ({"filters": [{"col": "total_quantity", "op": "contains", "value": "1"}]}, "operator invalid for the column kind"),
    ({"filters": [{"col": "status", "op": "eq", "value": "  "}]}, "blank value would be silently dropped"),
    ({"filters": [{"col": "expiry_date", "op": "before", "value": "not-a-date"}]}, "unparsable date"),
    ({"columns": ["zzz"]}, "unknown column"),
    ({"dataset": "risks"}, "dataset outside the Ava allowlist"),
    ({"frequency": "hourly"}, "bad frequency"),
    ({"format": "html"}, "bad format"),
    ({"name": "  "}, "no name"),
    ({"recipients": []}, "no recipients"),
])
def test_create_rejects_bad_input(env, over, why):
    with pytest.raises(S.ScheduleError):
        make(env, **over)


def test_create_requires_a_source(env):
    with pytest.raises(S.ScheduleError):
        S.create_schedule(env.db, env.admin, env.tid, {"name": "x", "recipients": ["a@b.co"]})


def test_create_forbidden_without_dataset_permission(env):
    with pytest.raises(S.ScheduleForbidden):
        make(env, user=env.viewer)


def test_saved_report_source_and_visibility(env):
    d = m.ReportDefinition(tenant_id=env.tid, slug="lic-active", name="Active licences", dataset="licenses",
                           created_by=env.admin.id, is_shared=False,
                           spec={"columns": ["publisher", "product"], "format": "html",
                                 "filters": [{"col": "status", "op": "eq", "value": "active"}]})
    env.db.add(d)
    env.db.commit()
    row = S.create_schedule(env.db, env.admin, env.tid, {"name": "From saved", "saved_report_slug": "lic-active",
                                                       "recipients": ["boss@acme.test"]})
    s = S.serialize(env.db, [row], env.admin.id)[0]
    assert row.saved_report_id == d.id and s["saved_report_name"] == "Active licences" and not s["saved_report_missing"]
    with pytest.raises(S.ScheduleError):                       # unknown slug
        S.create_schedule(env.db, env.admin, env.tid, {"name": "n", "saved_report_slug": "nope", "recipients": ["a@b.co"]})
    env.db.rollback()
    with pytest.raises(S.ScheduleError):                       # someone else's PRIVATE report
        S.create_schedule(env.db, env.other, env.tid, {"name": "n", "saved_report_slug": "lic-active", "recipients": ["a@b.co"]})
    env.db.rollback()
    d.is_shared = True
    env.db.commit()
    # A shared report is VISIBLE to other users, but sharing never grants the data behind it ...
    with pytest.raises(S.ScheduleForbidden):
        S.create_schedule(env.db, env.other, env.tid, {"name": "n", "saved_report_slug": "lic-active", "recipients": ["a@b.co"]})
    env.db.rollback()
    # ... once they hold the dataset's module permission they can schedule it.
    grant(env, env.other, "assets:asset_inventory:view")
    S.create_schedule(env.db, env.other, env.tid, {"name": "n", "saved_report_slug": "lic-active", "recipients": ["a@b.co"]})


# ── file generation (the builder's own export path) ─────────────────────────

def _spec(**over):
    spec = {"dataset": "licenses", "columns": ["publisher", "product"], "logic": "AND", "search": None, "sorts": [],
            "filters": [{"col": "status", "op": "eq", "value": "active"}], "source_label": None}
    spec.update(over)
    return spec


def test_generate_csv_xlsx_pdf(env):
    csv = S.generate_report_file(env.db, env.admin, _spec(), "csv", name="From saved", now=NOW)
    lines = csv.data.decode("utf-8").strip().splitlines()
    assert lines[0].strip() == "Publisher,Product" and len(lines) == 4          # 3 active rows + header
    assert csv.filename == "from-saved-2026-10-15.csv" and csv.mimetype == "text/csv"
    assert (csv.rows_total, csv.rows_exported) == (3, 3)

    import openpyxl
    xlsx = S.generate_report_file(env.db, env.admin, _spec(), "xlsx", name="x", now=NOW)
    ws = openpyxl.load_workbook(io.BytesIO(xlsx.data)).active
    assert ws.max_row == 4 and [c.value for c in ws[1]] == ["Publisher", "Product"]
    assert xlsx.filename.endswith(".xlsx") and "spreadsheetml" in xlsx.mimetype

    pdf = S.generate_report_file(env.db, env.admin, _spec(), "pdf", name="x", now=NOW)
    assert pdf.data[:4] == b"%PDF" and pdf.filename.endswith(".pdf")


def test_generate_rechecks_permission_and_refuses_unusable_filters(env):
    with pytest.raises(S.ScheduleForbidden):
        S.generate_report_file(env.db, env.viewer, _spec(), "csv", name="x", now=NOW)
    bad = _spec(filters=[{"col": "expiry_date", "op": "before", "value": "garbage"}])
    with pytest.raises(S.ScheduleError):                       # must NOT silently export every row
        S.generate_report_file(env.db, env.admin, bad, "csv", name="x", now=NOW)


# ── SMTP configuration order ────────────────────────────────────────────────

def test_smtp_resolution_order(env, monkeypatch):
    assert mailer.resolve_smtp(env.db, env.tid) is None
    assert mailer.email_status(env.db, env.tid) == {"configured": False, "source": None, "from_email": None}

    monkeypatch.setenv("SMTP_HOST", "smtp.env.test")
    monkeypatch.setenv("SMTP_USER", "envuser")
    monkeypatch.setenv("SMTP_PASSWORD", "envpw")
    monkeypatch.setenv("SMTP_PORT", "465")
    sm = mailer.resolve_smtp(env.db, env.tid)
    assert sm.source == "env" and sm.use_tls is False            # 465 -> implicit SSL, like workflow's send_email

    env.db.add(m.WorkflowEmailConfiguration(tenant_id=env.tid, config_name="primary", smtp_host="smtp.tenant.test",
                                            smtp_port=2525, smtp_username="t", smtp_password="tpw",
                                            from_email="noreply@tenant.test", from_name="Tenant", use_tls=True,
                                            is_active=True))
    env.db.commit()
    sm = mailer.resolve_smtp(env.db, env.tid)
    assert (sm.source, sm.host, sm.port, sm.from_name) == ("tenant", "smtp.tenant.test", 2525, "Tenant")

    env.db.query(m.WorkflowEmailConfiguration).update({"is_active": False})
    env.db.commit()
    assert mailer.resolve_smtp(env.db, env.tid).source == "env"  # inactive tenant config is ignored


# ── running: the email, and every way it can fail ───────────────────────────

def test_run_without_email_setup_records_the_failure(env):
    row = make(env)
    res = S.run_schedule(env.db, row, trigger="manual", actor="Alice Admin")
    assert res["ok"] is False and res["message"] == mailer.NOT_CONFIGURED_MESSAGE
    env.db.refresh(row)
    assert row.last_status == "error" and row.last_error == mailer.NOT_CONFIGURED_MESSAGE and row.last_run_at == NOW
    assert row.next_run == NOW + timedelta(hours=2)              # a manual run never moves the timed slot


def test_run_emails_the_report_as_an_attachment(env, smtp):
    row = make(env, name="Active <weekly>", message="Weekly <b>licence</b> review", recipients=["boss@acme.test", "ops@acme.test"])
    res = S.run_schedule(env.db, row, trigger="manual", actor="Alice Admin")
    assert res["ok"] is True and res["status"] == "sent" and res["rows"] == 3 and res["filename"].endswith(".csv")

    frm, to, raw = smtp.sent[-1]
    assert frm == "reports@env.test" and to == ["boss@acme.test", "ops@acme.test"]
    assert smtp.logins[-1] == ("envuser", "envpw")
    msg = email.message_from_string(raw)
    assert msg["Subject"] == "[Ava] Active <weekly> - 2026-10-15" and msg["To"] == "boss@acme.test, ops@acme.test"
    parts = list(msg.walk())
    attachments = [p for p in parts if p.get_filename()]
    assert len(attachments) == 1 and attachments[0].get_filename().endswith(".csv")
    assert attachments[0].get_payload(decode=True).decode("utf-8").splitlines()[0].strip() == "Publisher,Product,Status"
    html = [p for p in parts if p.get_content_type() == "text/html"][0].get_payload(decode=True).decode("utf-8")
    assert "&lt;weekly&gt;" in html and "<weekly>" not in html                   # user text is escaped
    assert "&lt;b&gt;licence&lt;/b&gt;" in html and "<b>licence" not in html
    assert "Software Licenses" in html and "Acme Corp" in html
    plain = [p for p in parts if p.get_content_type() == "text/plain"][0].get_payload(decode=True).decode("utf-8")
    assert "Sent on demand by Alice Admin." in plain
    env.db.refresh(row)
    assert row.last_status == "sent" and row.last_error is None


def test_run_smtp_problems_are_reported_not_raised(env, smtp):
    row = make(env, recipients=["boss@acme.test", "ops@acme.test"])
    smtp.fail_auth = True
    res = S.run_schedule(env.db, row, trigger="manual")
    assert res["ok"] is False and "username/password" in res["message"]
    smtp.fail_auth = False
    smtp.refuse = {"ops@acme.test": (550, b"no such user")}
    res = S.run_schedule(env.db, row, trigger="manual")
    assert res["ok"] is True and "ops@acme.test" in res["refused"] and "ops@acme.test" in res["message"]


def test_run_owner_and_source_problems_are_clear_errors(env, smtp):
    row = make(env)
    row.created_by = env.viewer.id                              # owner without the dataset permission
    env.db.commit()
    res = S.run_schedule(env.db, row, trigger="scheduled")
    assert res["ok"] is False and "permission" in res["message"].lower()

    env.viewer.is_active = False
    env.db.commit()
    res = S.run_schedule(env.db, row, trigger="scheduled")
    assert res["ok"] is False and "no longer has an active account" in res["message"]

    d = m.ReportDefinition(tenant_id=env.tid, slug="gone", name="Gone", dataset="licenses", created_by=env.admin.id,
                           spec={"columns": ["publisher"]})
    env.db.add(d)
    env.db.commit()
    srow = S.create_schedule(env.db, env.admin, env.tid, {"name": "Follows gone", "saved_report_slug": "gone",
                                                         "recipients": ["a@b.co"]})
    env.db.delete(d)
    env.db.commit()
    res = S.run_schedule(env.db, srow, trigger="manual")
    assert res["ok"] is False and "no longer exists" in res["message"]
    assert S.serialize(env.db, [srow], env.admin.id)[0]["saved_report_missing"] is True


def test_run_refuses_an_oversize_attachment(env, smtp, monkeypatch):
    big = S.GeneratedReport("x.csv", "text/csv", b"x" * (S.MAX_ATTACHMENT_BYTES + 1), 1, 1)
    monkeypatch.setattr(S, "generate_report_file", lambda *a, **k: big)
    res = S.run_schedule(env.db, make(env), trigger="manual")
    assert res["ok"] is False and "email limit" in res["message"] and not smtp.sent


# ── due schedules: claimed exactly once ─────────────────────────────────────

def test_claim_due_is_exactly_once(env):
    row = make(env, start_at=NOW + timedelta(hours=2))
    assert S.claim_due(env.db, NOW) == []                                  # not due yet
    later = NOW + timedelta(hours=3)
    assert S.claim_due(env.db, later) == [row.id]
    env.db.refresh(row)
    assert row.next_run > later                                            # advanced to the next slot
    assert S.claim_due(env.db, later) == []                                # a second tick wins nothing


def test_claim_due_loses_a_race_without_double_sending(env, monkeypatch):
    row = make(env, start_at=NOW + timedelta(hours=2))
    later = NOW + timedelta(hours=3)
    real = S._due_snapshot

    def snapshot_then_rival_claims(db, now):
        snap = real(db, now)                                               # we read next_run = T0 ...
        for sid, *_ in snap:                                               # ... then another worker claims it
            db.query(ScheduledReport).filter(ScheduledReport.id == sid).update(
                {"next_run": later + timedelta(days=1)}, synchronize_session=False)
            db.commit()
        return snap

    monkeypatch.setattr(S, "_due_snapshot", snapshot_then_rival_claims)
    assert S.claim_due(env.db, later) == []


def test_run_due_sends_what_it_claimed(env, smtp):
    row = make(env, start_at=NOW + timedelta(hours=2))
    later = NOW + timedelta(hours=3)
    out = S.run_due(env.db, later)
    assert out == {"claimed": 1, "sent": 1, "failed": 0} and len(smtp.sent) == 1
    env.db.refresh(row)
    assert row.last_status == "sent"
    assert S.run_due(env.db, later) == {"claimed": 0, "sent": 0, "failed": 0}
    assert len(smtp.sent) == 1


def test_pause_resume_and_edit(env):
    row = make(env)
    S.update_schedule(env.db, env.admin, row, {"active": False})
    assert row.active is False and row.next_run is None
    assert S.claim_due(env.db, NOW + timedelta(days=30)) == []
    S.update_schedule(env.db, env.admin, row, {"active": True})
    assert row.active is True and row.next_run > NOW
    S.update_schedule(env.db, env.admin, row, {"frequency": "monthly", "recipients": "x@y.co, z@y.co"})
    assert row.frequency == "monthly" and row.recipients == ["x@y.co", "z@y.co"]
    S.update_schedule(env.db, env.admin, row, {"name": "Renamed", "columns": ["publisher"]})
    assert row.name == "Renamed" and row.columns == ["publisher"]
    with pytest.raises(S.ScheduleError):
        S.update_schedule(env.db, env.admin, row, {"filters": [{"col": "status", "op": "gt", "value": "1"}]})
    env.db.rollback()


# ── HTTP layer ──────────────────────────────────────────────────────────────

def test_http_crud_access_control_and_run(env, smtp):
    pytest.importorskip("httpx")
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from grc.models import get_db
    from grc.modules.report_scheduler.router import router
    from grc.routers.auth_router import require_auth

    app = FastAPI()
    app.include_router(router)
    cur = {"user": env.admin}
    app.dependency_overrides[get_db] = lambda: env.db
    app.dependency_overrides[require_auth] = lambda: cur["user"]
    c = TestClient(app)

    j = c.get("/report-schedules").json()
    assert j["frequencies"] == ["daily", "weekly", "monthly"] and j["email"]["configured"] is True

    r = c.post("/report-schedules", json={"name": "HTTP made", "dataset": "licenses", "columns": ["publisher"],
                                          "filters": [{"col": "status", "op": "eq", "value": "active"}],
                                          "recipients": ["p@q.rs"], "frequency": "daily", "format": "pdf",
                                          "start_at": "2030-01-01T09:30:00Z"})
    assert r.status_code == 200 and r.json()["next_run"] == "2030-01-01T09:30:00Z"
    sid = r.json()["id"]

    r = c.post("/report-schedules", json={"name": "bad", "dataset": "licenses", "recipients": ["nope"]})
    assert r.status_code == 400 and "valid email" in r.json()["detail"]

    r = c.put(f"/report-schedules/{sid}", json={"active": False})
    assert r.json()["active"] is False and r.json()["next_run"] is None
    c.put(f"/report-schedules/{sid}", json={"active": True})

    r = c.post(f"/report-schedules/{sid}/run")
    j = r.json()
    assert r.status_code == 200 and j["ok"] is True and j["filename"].endswith(".pdf") and j["schedule"]["last_status"] == "sent"

    cur["user"] = env.other                                            # someone else's schedule: 404, not 403
    assert c.get(f"/report-schedules/{sid}").status_code == 404
    assert c.delete(f"/report-schedules/{sid}").status_code == 404
    assert c.get("/report-schedules").json()["schedules"] == []

    cur["user"] = env.viewer                                           # no dataset permission
    r = c.post("/report-schedules", json={"name": "nope", "dataset": "licenses", "recipients": ["a@b.co"]})
    assert r.status_code == 403

    cur["user"] = env.admin                                            # administrators see the whole workspace
    assert any(s["id"] == sid for s in c.get("/report-schedules").json()["schedules"])
    assert c.delete(f"/report-schedules/{sid}").json()["deleted"] is True
    assert c.get(f"/report-schedules/{sid}").status_code == 404
    assert c.post("/report-schedules/999999/run").status_code == 404


# ── Celery fan-out ──────────────────────────────────────────────────────────

def test_celery_fan_out_claims_and_dispatches_once(env, smtp, monkeypatch):
    pytest.importorskip("celery")
    import grc.db as gdb
    from grc.modules.report_scheduler import tasks

    monkeypatch.setattr(gdb, "open_tenant_session", lambda slug: env.Session())

    class Master:
        def query(self, *a): return self
        def filter(self, *a): return self
        def all(self): return [SimpleNamespace(slug="acme")]
        def close(self): pass

    monkeypatch.setattr(gdb, "MasterSession", lambda: Master())

    due = make(env, name="Due")
    later = make(env, name="Later", start_at=NOW + timedelta(days=3))
    due_id, later_id = due.id, later.id
    env.db.execute(ScheduledReport.__table__.update().where(ScheduledReport.id == due_id)
                   .values(next_run=datetime.utcnow() - timedelta(minutes=3)))
    env.db.commit()
    monkeypatch.setattr(S, "_utcnow", datetime.utcnow)                 # the task path uses the real clock

    dispatched = []
    monkeypatch.setattr(tasks.run_scheduled_report_task, "delay",
                        lambda **kw: dispatched.append(kw) or tasks.run_scheduled_report_task.run(**kw))
    assert tasks.scheduled_reports_fan_out.run() == {"status": "ok", "dispatched": 1}
    assert dispatched == [{"tenant_slug": "acme", "schedule_id": due_id}] and len(smtp.sent) == 1
    s2 = env.Session()
    assert s2.get(ScheduledReport, due_id).last_status == "sent"
    assert s2.get(ScheduledReport, later_id).last_status is None
    assert tasks.scheduled_reports_fan_out.run()["dispatched"] == 0 and len(smtp.sent) == 1
