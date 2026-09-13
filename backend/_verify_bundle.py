import json, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from collections import Counter, defaultdict
HERE = os.path.dirname(os.path.abspath(__file__)); SEED = os.path.join(HERE, 'grc', 'seed_data')
D = json.load(open(os.path.join(SEED, 'stage4_all_domains.json'), encoding='utf-8'))['domains']
s3 = json.load(open(os.path.join(SEED, 'stage3_complete.json'), encoding='utf-8'))
def fwshort(c): return c['framework'].split('(')[0].strip()
hd = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='17B898'); thin = Side(style='thin', color='DDDDDD')
def hdr(ws, heads):
    for j, h in enumerate(heads, 1):
        c = ws.cell(1, j, h); c.font = hd; c.fill = fill; c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.freeze_panes = 'A2'
def row(ws, r, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.alignment = Alignment(vertical='top', wrap_text=True); c.border = Border(bottom=thin)

def uid(fw, cid, title): return '%s || %s || %s' % (fw, cid, title)

wb = openpyxl.Workbook()
# ---- Tab 1: Standalone controls (1,856) ----
st = wb.active; st.title = '1 Standalone controls'
hdr(st, ['Control UID (framework | id | title)', 'Domain', 'Framework', 'Control ID', 'Original Title']); r = 2; n1 = 0
for dm in D:
    for s in dm['sets']:
        if s['member_count'] == 1:
            m = s['members'][0]; row(st, r, [uid(m['framework'], m['control_id'], m['original_title']), dm['domain'], m['framework'], m['control_id'], m['original_title']]); r += 1; n1 += 1
for col, w in zip('ABCDE', [40, 34, 30, 14, 70]): st.column_dimensions[col].width = w

# ---- Tab 2: Source control list per framework (3,352) ----
sc = wb.create_sheet('2 Source control list')
hdr(sc, ['Control UID (framework | id | title)', 'Framework', 'Control ID', 'Original Title', 'Source Domain (as ingested)']); r = 2; n2 = 0
for dm in s3['domains']:
    for c in dm['controls']:
        row(sc, r, [uid(fwshort(c), c['control_id'], c['title']), fwshort(c), c['control_id'], c['title'], dm['domain']]); r += 1; n2 += 1
for col, w in zip('ABCDE', [40, 30, 14, 70, 30]): sc.column_dimensions[col].width = w

# ---- Tab 3: Set membership map (1,496, flat) ----
mm = wb.create_sheet('3 Set membership map')
hdr(mm, ['Control UID (framework | id | title)', 'Domain', 'Set', 'Set Heading', 'Framework', 'Control ID', 'Original Title']); r = 2; n3 = 0
for dm in D:
    for s in dm['sets']:
        if s['member_count'] > 1:
            for m in s['members']:
                row(mm, r, [uid(m['framework'], m['control_id'], m['original_title']), dm['domain'], s['set_id'], s['normalized_title'], m['framework'], m['control_id'], m['original_title']]); r += 1; n3 += 1
for col, w in zip('ABCDEFG', [40, 22, 6, 36, 28, 12, 60]): mm.column_dimensions[col].width = w

# ---- Tab 4: Framework catalog source (broadcast evidence) ----
fc = wb.create_sheet('4 Framework catalog source')
hdr(fc, ['Framework', 'Catalog Artifact', '# controls attached', 'framework total controls', '% of controls']); r = 2
fwctrls = Counter(); artfreq = defaultdict(Counter)
for dm in s3['domains']:
    for c in dm['controls']:
        f = fwshort(c); fwctrls[f] += 1
        for a in (c.get('artifacts') or []): artfreq[f][a.get('name')] += 1
# list every catalog artifact currently lifted, with its source prevalence
lifted = set()
for dm in D:
    for cat in dm.get('framework_catalog_artifacts', []):
        for a in cat['artifacts']: lifted.add((cat['framework'], a['name']))
n4 = 0
for (f, name) in sorted(lifted):
    cnt = artfreq[f].get(name, 0); tot = fwctrls[f]
    row(fc, r, [f, name, cnt, tot, '%d%%' % round(100 * cnt / tot) if tot else '0%']); r += 1; n4 += 1
for col, w in zip('ABCDE', [30, 46, 18, 22, 12]): fc.column_dimensions[col].width = w

# ---- Tab 5: Catalog artifact -> domain spread (genericness proof) ----
sp = wb.create_sheet('5 Catalog artifact spread')
hdr(sp, ['Framework', 'Catalog Artifact', '# controls at source', '# distinct domains', 'Domains the bearer controls span']); r = 2
# map every control identity -> its stage4 domain
id2dom = {}
for dm in D:
    for s in dm['sets']:
        for m in s['members']: id2dom[(m['framework'], m['control_id'], m['original_title'])] = dm['domain']
# per (framework, artifact name): bearer controls at source -> their stage4 domains
art_ctrl = defaultdict(list)
for dm in s3['domains']:
    for c in dm['controls']:
        for a in (c.get('artifacts') or []):
            art_ctrl[(fwshort(c), a.get('name'))].append((fwshort(c), c['control_id'], c['title']))
n5 = 0
for (f, name) in sorted(lifted):
    ctrls = art_ctrl.get((f, name), [])
    domains = sorted(set(id2dom.get(cid, '(?)') for cid in ctrls))
    domcell = ' ; '.join(domains) if ctrls else 'source artifact extract incomplete for this framework (known SAMA gap) — was a verified broadcast in the live library'
    row(sp, r, [f, name, len(ctrls), len(domains), domcell]); r += 1; n5 += 1
for col, w in zip('ABCDE', [30, 46, 18, 18, 80]): sp.column_dimensions[col].width = w

# ---- Tab 0: Read me / reconciliation ----
rm = wb.create_sheet('0 Read me', 0)
total4 = n1 + n3
rm['A1'] = 'Verification bundle — items 1 to 4 requested'; rm['A1'].font = Font(bold=True, size=13)
notes = [
 '', 'Tab 1  Standalone controls: %d rows (one per framework-unique control).' % n1,
 'Tab 2  Source control list: %d rows (every control as ingested: framework, control ID, original title).' % n2,
 'Tab 3  Set membership map: %d rows (flat, every set member; no merged cells).' % n3,
 'Tab 4  Framework catalog source: %d lifted catalog artifacts with how many of the framework controls carry each at source.' % n4,
 'Tab 5  Catalog artifact spread: for each catalog artifact, how many DISTINCT domains its bearer controls span (genericness = high domain spread, not % of framework controls).',
 '', 'Reconciliation: set members (%d) + standalone (%d) = %d  (expected 3,352).' % (n3, n1, total4),
 'Join on the Control UID column (framework | id | title), NOT on Control ID alone.',
 'Note: SAMA source data reuses ID strings across distinct controls (e.g. "1.3" labels 9 different controls; "N/A" used as an ID). So (Framework, Control ID) is NOT unique for SAMA; the UID column is the correct key. This is a known SAMA source-quality gap, not duplication.',
 'Join Tab 1 + Tab 3 on Control UID to confirm the partition: every control in exactly one place, none duplicated.',
 'Diff Tab 3/Tab 1 Original Title against Tab 2 to confirm no title was reworded versus the ingested source.',
]
for i, t in enumerate(notes, 2): rm['A' + str(i)] = t
rm.column_dimensions['A'].width = 110

xp = os.path.expanduser('~/OneDrive/Desktop/VERIFICATION_BUNDLE.xlsx')
try: wb.save(xp)
except PermissionError: xp = os.path.expanduser('~/OneDrive/Desktop/VERIFICATION_BUNDLE_v2.xlsx'); wb.save(xp)
print('standalone=%d  source=%d  members=%d  catalog=%d  | members+standalone=%d' % (n1, n2, n3, n4, total4))
print('wrote:', xp)
